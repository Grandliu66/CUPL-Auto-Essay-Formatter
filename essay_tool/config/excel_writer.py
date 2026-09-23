"""把内存里的 PresetConfig 写成一份 presets/*.xlsx——Excel 这条路径现在是
"边缘产品"：默认存盘走 JSON（见 `preset_io.py`），这里只管"用户想拿 Excel
可视化看一眼/改一改"这条边缘用途，所以只保证读写正确、格式跟
`tools/build_presets.py` 生成的表兼容（能被 `excel_io.load_preset` 读回来），
不追求视觉上跟那份手工表一模一样。

单元格里写的字符串复用 `serialize.py` 那套转换（跟 JSON 存盘同一份"人读得懂
的写法"），两条路径的文字表示保持一致。
"""

from __future__ import annotations

import pathlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from essay_tool.config import serialize, units
from essay_tool.config.excel_io import SHEET_NAME
from essay_tool.config.schema import COVER_FORM_NAMES, HEADING_LABEL_PARTS, PresetConfig

# "插入文字"是内容占位语的四个部件（题目本身是内容，没有标准答案可填）
_PLACEHOLDER_PARTS = frozenset({"题目", "副标题", "英文题目", "英文副标题"})

MAX_COLS = 19

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="center", wrap_text=True)
CENTER = Alignment(vertical="center", horizontal="center", wrap_text=True)

YN = ["是", "否"]
ALIGN = list(units.ALIGNMENTS)
LINE_RULE = list(units.LINE_RULES)
FONT_SIZES = list(units.CHINESE_FONT_SIZES)
INDENT_UNITS = ["字符", "厘米", "-"]
SPACING_UNITS = ["磅"]
PAPER_SIZES = ["A4", "A3", "B5", "自定义"]
DIRECTIONS = ["纵向", "横向"]
PAGE_NUMBER_FORMATS = list(units.PAGE_NUMBER_FORMATS)
_PAGE_NUMBER_CODE_TO_DISPLAY = {v: k for k, v in units.PAGE_NUMBER_FORMATS.items()}

# 「清理规则」保留的字符格式：内部代码（w:rPr 子元素名）->中文标签，跟
# excel_io._KEEPABLE_RUN_FORMATS 是同一份词汇的反查（各自维护一份，避免
# 这个"边缘产品"跟主读取路径产生私有成员耦合）。
_RUN_FORMAT_LABELS = {
    "b": "加粗", "i": "斜体", "u": "下划线", "vertAlign": "上标下标",
    "strike": "删除线", "dstrike": "双删除线", "color": "颜色",
    "highlight": "高亮", "vanish": "隐藏文字", "em": "着重号", "spacing": "字符间距",
}
_DEFAULT_KEEP_RUN_FORMATTING_LABEL = "加粗、斜体、下划线、上标下标"

COL_WIDTHS = [10, 16, 12, 16, 8, 6, 10, 10, 8, 10, 8, 8, 8, 8, 8, 12, 8, 10, 22]

NOTE_TEXT = (
    "说明\n"
    "1. 该Excel是可供工具直接读取的配置文件来源，所有的填写都必须按格式进行，否则将无法为工具正确读取。\n"
    '2. "备注"栏仅为分享和自我提醒之用，无法为工具理解。\n'
    "3. 本文件由「保存当前设置为新配置」/「导出为Excel」导出，仅供人工查看和可视化修改；"
    "程序默认从 JSON 配置文件读写，这份 Excel 改完后需要重新「从Excel导入配置文件」才会生效。"
)


def _intro_note(ws, row: int, text: str) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=MAX_COLS)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(size=10, color="595959")
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 62
    return row + 1


def _section_title(ws, row: int, text: str) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=MAX_COLS)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(bold=True, size=13, color="FFFFFF")
    cell.fill = TITLE_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 24
    return row + 1


def _header_row(ws, row: int, headers: list[str]) -> int:
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    return row + 1


def _data_rows(ws, row: int, rows: list[list], center_cols: set[int]) -> int:
    for offset, data in enumerate(rows):
        r = row + offset
        for c, val in enumerate(data, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = BORDER
            cell.alignment = CENTER if c in center_cols else WRAP
    return row + len(rows)


def _add_block(ws, row: int, title: str, headers: list[str], rows: list[list],
                center_cols: set[int] = frozenset()) -> tuple[int, int]:
    row = _section_title(ws, row, title)
    row = _header_row(ws, row, headers)
    data_start = row
    row = _data_rows(ws, row, rows, center_cols)
    return data_start, row


def _add_validation(ws, col: int, row_start: int, row_end: int, formula1: str,
                    allow_blank: bool = False) -> None:
    if row_start >= row_end:
        return
    dv = DataValidation(type="list", formula1=formula1, allow_blank=allow_blank)
    ws.add_data_validation(dv)
    letter = get_column_letter(col)
    dv.add(f"{letter}{row_start}:{letter}{row_end - 1}")


def _add_dropdown(ws, col: int, row_start: int, row_end: int, options: list[str],
                  allow_blank: bool = False) -> None:
    _add_validation(ws, col, row_start, row_end, '"' + ",".join(options) + '"',
                    allow_blank=allow_blank)


def _add_dropdown_range(ws, col: int, row_start: int, row_end: int, ref: str,
                        allow_blank: bool = False) -> None:
    _add_validation(ws, col, row_start, row_end, ref, allow_blank=allow_blank)


def _keep_run_formatting_label(cfg: PresetConfig) -> str:
    if cfg.keep_run_formatting is None:
        return _DEFAULT_KEEP_RUN_FORMATTING_LABEL
    if not cfg.keep_run_formatting:
        return "无"
    labels = [_RUN_FORMAT_LABELS.get(code, code) for code in cfg.keep_run_formatting]
    return "、".join(labels)


def _write(ws, cfg: PresetConfig, page_number_format_range: str) -> None:
    row = 1
    row = _intro_note(ws, row, NOTE_TEXT)
    row += 1

    # ------------------------------------------------------------ 页面设置
    p = serialize.page_setup_to_dict(cfg.page)
    is_named = p["paper_size"] in units.NAMED_PAPER_SIZES
    headers = ["项目", "值", "备注"]
    rows = [
        ["纸张", p["paper_size"], ""],
        ["自定义宽（厘米）", "" if is_named else p["page_width_cm"], ""],
        ["自定义高（厘米）", "" if is_named else p["page_height_cm"], ""],
        ["纸张方向", p["orientation"], ""],
        ["上边距（厘米）", p["margin_top_cm"], ""],
        ["下边距（厘米）", p["margin_bottom_cm"], ""],
        ["左边距（厘米）", p["margin_left_cm"], ""],
        ["右边距（厘米）", p["margin_right_cm"], ""],
        ["装订线（厘米）", p["gutter_cm"], ""],
        ["页眉距边（厘米）", p["header_distance_cm"], ""],
        ["页脚距边（厘米）", p["footer_distance_cm"], ""],
    ]
    data_start, data_end = _add_block(ws, row, "页面设置", headers, rows, center_cols={2, 3})
    _add_dropdown(ws, 2, data_start, data_start + 1, PAPER_SIZES)
    _add_dropdown(ws, 2, data_start + 3, data_start + 4, DIRECTIONS)
    row += 1 + 1 + len(rows) + 1

    # ----------------------------------------------------------------- 样式
    headers = [
        "编号", "名称", "中文字体", "西文字体", "字号", "加粗", "对齐",
        "首行缩进单位", "首行缩进值", "左缩进单位", "左缩进值",
        "段前单位", "段前值", "段后单位", "段后值",
        "行距类型", "行距值", "段前分页", "备注",
    ]
    rows = []
    for spec in cfg.styles.values():
        sd = serialize.style_spec_to_dict(spec)
        rows.append([
            sd["id"], sd["name"], sd["font_east_asian"], sd["font_ascii"], sd["size"],
            sd["bold"], sd["alignment"], sd["first_line_indent_unit"], sd["first_line_indent_value"],
            sd["left_indent_unit"], sd["left_indent_value"], "磅", sd["space_before"],
            "磅", sd["space_after"], sd["line_rule"], sd["line_value"], sd["page_break_before"], "",
        ])
    data_start, data_end = _add_block(
        ws, row, "样式", headers, rows,
        center_cols={1, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18},
    )
    _add_dropdown(ws, 5, data_start, data_end, FONT_SIZES)
    _add_dropdown(ws, 6, data_start, data_end, YN)
    _add_dropdown(ws, 7, data_start, data_end, ALIGN)
    _add_dropdown(ws, 8, data_start, data_end, INDENT_UNITS)
    _add_dropdown(ws, 10, data_start, data_end, INDENT_UNITS)
    _add_dropdown(ws, 12, data_start, data_end, SPACING_UNITS)
    _add_dropdown(ws, 14, data_start, data_end, SPACING_UNITS)
    _add_dropdown(ws, 16, data_start, data_end, LINE_RULE)
    _add_dropdown(ws, 18, data_start, data_end, YN)
    row += 1 + 1 + len(rows) + 1

    # --------------------------------------------------------------- 页眉页脚
    headers = ["节", "页眉文字", "启用页码", "页码编码", "前缀", "后缀", "备注"]
    rows = []
    for r_ in cfg.header_footer:
        hd = serialize.header_footer_rule_to_dict(r_)
        # 页码编码：`""`→空（＝没配）；认得的一律转成中文显示值；**认不出来的
        # 写原值**（审计 2026-09-15 核实）：以前认不出来就写 `""`，读回来变
        # None——一个工具不认识的编码（比如从别处拷来的 `decimalZero`）会在
        # “导出成 Excel”这一步静默消失。宁可在表里看着眼生，也不能丢。
        code = hd["page_number_format"]
        fmt_display = (_PAGE_NUMBER_CODE_TO_DISPLAY.get(code, code) if code else "")
        rows.append([hd["section_name"], hd["header_text"], hd["page_number_enabled"],
                     fmt_display, hd["page_number_prefix"], hd["page_number_suffix"], ""])
    data_start, data_end = _add_block(ws, row, "页眉页脚", headers, rows, center_cols={3, 4, 5, 6})
    _add_dropdown(ws, 3, data_start, data_end, YN)
    _add_dropdown_range(ws, 4, data_start, data_end, page_number_format_range)
    row += 1 + 1 + len(rows) + 1

    # ---------------------------------------------------------------- 页眉线
    headers = ["项目", "值", "备注"]
    rows = [
        ["页眉文字下画横线", "要" if cfg.header_bottom_line else "不要", ""],
        ["页眉线左伸出（厘米）", cfg.header_line_extend_left_cm, "线向正文栏外伸出多少；0＝与正文栏等宽"],
        ["页眉线右伸出（厘米）", cfg.header_line_extend_right_cm, "同上"],
    ]
    data_start, data_end = _add_block(ws, row, "页眉线", headers, rows, center_cols={2, 3})
    _add_dropdown(ws, 2, data_start, data_start + 1, ["要", "不要"])
    row += 1 + 1 + len(rows) + 1

    # -------------------------------------------------------------------- 引注
    cit = serialize.citation_to_dict(cfg.citation) or {}
    headers = ["项目", "值", "备注"]
    rows = [
        ["脚注序号样式", cit.get("footnote_marker_style", "①②③"), ""],
        ["脚注每页重新编号", cit.get("footnote_restart_each_page", "是"), ""],
        ["脚注序号字体", cit.get("footnote_marker_font", "正文(非上标)"), ""],
        ["脚注序号后空白", cit.get("footnote_marker_trailing_space", "半角空格"), ""],
        ["参考文献编号样式", cit.get("reference_numbering_style", "[1][2][3]"), ""],
    ]
    data_start, data_end = _add_block(ws, row, "引注", headers, rows, center_cols={2, 3})
    # ⚠️ 以前这里给的是 "1)2)3)"／"[1][2][3]"——`footnotes._NUMBER_FORMATS`
    # 根本不认这两种写法，选中之后引擎会当"认不出来"静默退回默认值
    # （2026-09-22 审计揪出的真 bug）。这一格是"脚注序号样式"，跟下面"参考
    # 文献编号样式"是两套完全不同的取值表，不能共用同一批选项；改成
    # `_NUMBER_FORMATS` 真正认识的写法。
    _add_dropdown(ws, 2, data_start, data_start + 1, ["①②③", "1.2.3.", "123"])
    _add_dropdown(ws, 2, data_start + 1, data_start + 2, YN)
    _add_dropdown(ws, 2, data_start + 2, data_start + 3, ["正文(非上标)", "上标"])
    _add_dropdown(ws, 2, data_start + 3, data_start + 4, ["半角空格", "全角空格", "不空格"])
    _add_dropdown(ws, 2, data_start + 4, data_start + 5, ["[1][2][3]", "1.2.3.", "(1)(2)(3)"])
    row += 1 + 1 + len(rows) + 1

    # --------------------------------------------------------------- 清理规则
    headers = ["项目", "值", "备注"]
    rows = [["保留的字符格式", _keep_run_formatting_label(cfg), ""],
            ["清除手打分页符", "是" if cfg.clear_page_breaks else "否", "封面范围除外；分节符不算"],
            ["清除空行", "是" if cfg.clear_blank_lines else "否", "封面范围除外；只清确实为空的段落"]]
    data_start, data_end = _add_block(ws, row, "清理规则", headers, rows, center_cols={2, 3})
    _add_dropdown(
        ws, 2, data_start, data_start + 1,
        ["加粗、斜体、下划线、上标下标", "加粗、斜体、下划线", "加粗、斜体", "无"],
    )
    _add_dropdown(ws, 2, data_start + 1, data_start + 3, YN)
    row += 1 + 1 + len(rows) + 1

    # ---------------------------------------------------------------- 封面
    # 三行：悬挂缩进、要不要封面、用哪种封面。后两行 2026-09-15 补（审计核实
    # 它们在 JSON 预设里能存、在 Excel 预英里却压根没有对应的格子 → 存成
    # Excel 就丢了，课程考核那份预设的"使用课程考核封面"也就传不过去）。
    # 留空＝预设里没写（旧预设/内置毕业论文那份就是这种），套用时保持界面
    # 上当前的选择不改。
    headers = ["项目", "值", "备注"]
    rows = [
        ["超长值悬挂缩进", "要" if cfg.cover_hanging_indent else "不要", ""],
        ["是否插入封面", "" if cfg.cover_enabled is None else ("要" if cfg.cover_enabled else "不要"),
         "留空＝套用这个预设时不改封面开关"],
        ["封面形态", cfg.cover_form or "",
         "留空＝套用时保持界面上当前选的形态（可填" + "/".join(COVER_FORM_NAMES) + "）"],
    ]
    data_start, data_end = _add_block(ws, row, "封面", headers, rows, center_cols={2, 3})
    _add_dropdown(ws, 2, data_start, data_start + 1, ["要", "不要"])
    _add_dropdown(ws, 2, data_start + 1, data_start + 2, ["要", "不要"], allow_blank=True)
    _add_dropdown(ws, 2, data_start + 2, data_start + 3, list(COVER_FORM_NAMES), allow_blank=True)
    row += 1 + 1 + len(rows) + 1

    # ------------------------------------------------------------ 结构部件
    # "插入文字"＝标签写法（缺失时插入的文字，也是识别到后改写成的写法）：可以
    # 自定义，空＝保持作者原文。题目/副标题/英文题目/英文副标题的这一格是"内容
    # 占位语"，工具行为、不许自定义，恒为空。引言/结论排在最后两行：它们没有
    # "是否启用"（正文里的一级标题，位置即身份），只有写法一格。
    # "标签前另起一页"（2026-09-20）：是/否＝这个部件的标签前要/不要另起一页（覆盖样式表
    # 里"段前分页"）；**空＝跟随样式**；"—"＝这个部件不是标签，没有这一项。
    def _pb(name: str) -> str:
        if name in cfg.label_page_break:
            return "是" if cfg.label_page_break[name] else "否"
        return ""

    headers = ["部件", "是否启用", "插入文字", "标签前另起一页", "备注"]
    rows = []
    for name, on in cfg.structure_required.items():
        if name in _PLACEHOLDER_PARTS:      # 占位语是工具行为，这一格恒空（引擎也不读）
            rows.append([name, "是" if on else "否", "", "—", "占位语固定"])
        else:
            rows.append([name, "是" if on else "否", cfg.structure_insert_text.get(name, ""),
                         _pb(name), ""])
    n_switchable = len(rows)
    for name in HEADING_LABEL_PARTS:
        rows.append([name, "—", cfg.structure_insert_text.get(name, ""), _pb(name),
                     "无开关，空＝原文"])
    data_start, data_end = _add_block(ws, row, "结构部件", headers, rows, center_cols={2, 4})
    _add_dropdown(ws, 2, data_start, data_start + n_switchable, YN)
    _add_dropdown(ws, 4, data_start, data_end, YN, allow_blank=True)
    row += 1 + 1 + len(rows) + 1

    for c, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w


def write_preset(cfg: PresetConfig, out_path: str | pathlib.Path) -> pathlib.Path:
    """把 cfg 写成一份 presets 风格的 xlsx，返回写出的路径。"""
    out_path = pathlib.Path(out_path)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    options_ws = wb.create_sheet("选项")
    for i, v in enumerate(PAGE_NUMBER_FORMATS, start=1):
        options_ws.cell(row=i, column=1, value=v)
    options_ws.sheet_state = "hidden"
    page_number_format_range = f"选项!$A$1:$A${len(PAGE_NUMBER_FORMATS)}"

    _write(ws, cfg, page_number_format_range)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
