"""⚠️ 已停用（2026-09-16）：从《工具操作逻辑写作表.xlsx》生成
`essay_tool/gui/operation_logic_text.py`。

用户 2026-09-16 裁决"停掉生成链"：查阅窗口里那份《工具操作逻辑》正文**已改成
手工维护**（直接改 `essay_tool/gui/operation_logic_text.py`），工作簿留在
`reference/` 当历史存档。这个脚本还在，是为了"万一想回到表格当源头"，但**直接
跑它会把窗口里的手工修改全部盖回去**，所以加了 `--force` 闸门。

跑法（真要用）：.venv\\Scripts\\python.exe tools\\make_operation_logic_text.py --force

映射规则（停用前的老规矩，留着备查）：

- 工作表「工具操作逻辑」：A 列章标题、B 列节标题（二者按“1.”“1.1”的写法直接
  成为窗口左侧树的层级）、C 列正文；C 列里形如「（sheet2的表）」的整行会被替换成
  对应工作表的真表格（sheetN 指工作簿里第 N 张表）。
- 表格：第 1 行作表头，其余为数据行；表头为空但该列有数据的列，按项目里的叫法
  补名（Sheet1 第 3 列＝「子项」）；整列没有数据的列不输出；纵向合并的单元格把值
  填到合并区的每一行；单元格内部的换行写成字面量 \\\\n（查阅窗口会把它显示成空格）。
- 某节正文还空着时，用 EMPTY_SECTION_FALLBACK 里的兜底引导语（并照常插表），
  工作簿补上正文之后即以工作簿为准。
"""

from __future__ import annotations

import pathlib
import re

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "reference" / "工具操作逻辑写作表.xlsx"
OUT_MODULE = ROOT / "essay_tool" / "gui" / "operation_logic_text.py"
OUTLINE_SHEET = "工具操作逻辑"

# 正文里的表格占位符：「（sheet2的表）」「（sheet3）的表」
PLACEHOLDER_RE = re.compile(r"^[（(]?\s*sheet\s*(\d+)\s*[）)]?\s*的\s*表\s*[）)]?$", re.I)

# 表头为空、但该列有数据时补的名字（沿用 rebuild_outline_sheet2.py 里的叫法）
COLUMN_NAMES = {"Sheet1": {3: "子项"}}

# 正文尚未写入工作簿的节：兜底引导语 + 要插进去的工作表序号（sheet3＝第 3 张表）
EMPTY_SECTION_FALLBACK: dict[str, tuple[str, int]] = {
    "6.2 其他告警": (
        "除识别告警外，工具在运行中报出的下列提示亦写进日志区。这些提示与识别告警一样，"
        "均不弹出中断性的对话框（程序崩溃另有专门的弹窗，见 6.3）：能够继续进行的照常进行，"
        "备份或清理失败不影响导入与导出；不能继续的则停于该步，例如成品正被 Word 占用时"
        "导出失败，不生成成品。各条提示的触发条件与原文如下表所示。",
        4,
    ),
}


def cell_text(ws, row: int, col: int, fill: dict[tuple[int, int], object]) -> str:
    """取一个格子的文字：合并区已填过值的用填的值；竖线换成全角以免破坏表格。"""
    value = fill.get((row, col), ws.cell(row=row, column=col).value)
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").strip()
    text = text.replace("|", "／")
    return text.replace("\n", "\\n")


def render_table(ws, sheet_name: str) -> list[str]:
    """把一张工作表渲染成 Markdown 管道表。"""
    fill: dict[tuple[int, int], object] = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col == rng.max_col and rng.max_row > rng.min_row:
            top = ws.cell(row=rng.min_row, column=rng.min_col).value
            for row in range(rng.min_row, rng.max_row + 1):
                fill[(row, rng.min_col)] = top

    ncols = ws.max_column
    heads = []
    for col in range(1, ncols + 1):
        head = cell_text(ws, 1, col, fill)
        heads.append(head or COLUMN_NAMES.get(sheet_name, {}).get(col, ""))
    rows = [[cell_text(ws, r, c, fill) for c in range(1, ncols + 1)]
            for r in range(2, ws.max_row + 1)]

    keep = [i for i in range(ncols) if any(row[i] for row in rows)]
    heads = [heads[i] for i in keep]
    rows = [[row[i] for i in keep] for row in rows]

    sep = [":--:" if head == "序号" else "---" for head in heads]
    lines = ["| " + " | ".join(heads) + " |",
             "| " + " | ".join(sep) + " |"]
    lines += ["| " + " | ".join(row) + " |" for row in rows]
    print(f"  表「{sheet_name}」：{len(rows)} 行 x {len(heads)} 列"
          f"（丢弃 {ncols - len(heads)} 个空列）")
    return lines


def main() -> None:
    book = openpyxl.load_workbook(WORKBOOK)
    sheets = book.sheetnames
    ws = book[OUTLINE_SHEET]

    out: list[str] = []
    troubles: list[str] = []
    tables: set[str] = set()

    def blank() -> None:
        if out and out[-1] != "":
            out.append("")

    def table_for(index: int, where: str) -> None:
        if not 1 <= index <= len(sheets):
            troubles.append(f"{where}：工作簿里没有第 {index} 张表")
            return
        name = sheets[index - 1]
        out.extend(render_table(book[name], name))
        tables.add(name)

    for row in range(2, ws.max_row + 1):
        chapter = str(ws.cell(row=row, column=1).value or "").strip()
        section = str(ws.cell(row=row, column=2).value or "").strip()
        body = str(ws.cell(row=row, column=3).value or "")

        if chapter:
            blank()
            out.append(chapter)
        if section:
            blank()
            out.append(section)

        lines = [line.strip() for line in body.split("\n")] if body.strip() else []
        inserted = False
        for line in lines:
            if not line:
                continue
            match = PLACEHOLDER_RE.match(line)
            if match:
                table_for(int(match.group(1)), f"{section or chapter} 的占位符 {line}")
                inserted = True
                blank()
            else:
                out.append(line)
        if inserted:
            continue

        if not lines and section:
            fallback = EMPTY_SECTION_FALLBACK.get(section)
            if fallback is None:
                troubles.append(f"{section}：正文还空着，生成出来的这一节只有标题")
                continue
            lead, index = fallback
            print(f"  {section}：工作簿里正文为空，使用兜底引导语 + 第 {index} 张表")
            out.append(lead)
            table_for(index, section)
            blank()

    text = "\n".join(out).strip("\n") + "\n"
    if '"""' in text or text.rstrip("\n").endswith("\\"):
        troubles.append("正文里出现三引号、或以反斜杠结尾，生成模块会不合法")

    header = (
        'r"""《工具操作逻辑》正文（由 tools/make_operation_logic_text.py 生成，勿手工编辑）。\n'
        "\n"
        f"来源：{WORKBOOK.relative_to(ROOT)}\n"
        "  * 「工具操作逻辑」表：章、节、正文（正文里的（sheetN的表）会换成对应的真表格）\n"
        f"  * 当前工作簿的表顺序：{('、'.join(sheets))}\n"
        "\n"
        "重新生成：.venv\\Scripts\\python.exe tools\\make_operation_logic_text.py\n"
        '"""\n\n'
    )
    OUT_MODULE.write_text(header + 'OPERATION_LOGIC_TEXT = r"""\n'
                          + text + '"""\n', encoding="utf-8")

    print(f"\n写出 {OUT_MODULE.relative_to(ROOT)}：{len(text.splitlines())} 行、"
          f"{len(text)} 字、插入了 {len(tables)} 张表（{'、'.join(sorted(tables))}）")
    if troubles:
        print("⚠ 需要注意：")
        for item in troubles:
            print(f"   {item}")


if __name__ == "__main__":
    import sys

    if "--force" not in sys.argv:
        print(
            "⚠️ 这个脚本已停用（2026-09-16）：《工具操作逻辑》正文改成手工维护，"
            "源头是 essay_tool\\gui\\operation_logic_text.py。\n"
            "   跑它会把手改的内容全部盖回去；确实要跑就加 --force。"
        )
        sys.exit(2)
    main()
