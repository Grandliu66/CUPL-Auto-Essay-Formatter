"""空行位置扫描：在文档的每个位置分别插入 1/2 个空行，两条导出路径（自动/结构页）各导两遍，
占位语不许重复、回炉数量必须稳定（2026-09-20 抓到过「标题后有空行→每导一遍多一条副标题占位语」）。

跑法：.venv\\Scripts\\python.exe tools\\check_blank_sweep_20260920.py    （退出码 0＝全部通过）
"""

import sys, pathlib, tempfile, dataclasses, collections
sys.stdout.reconfigure(encoding="utf-8")
ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from docx import Document
from essay_tool.engine import placeholders as ph
from essay_tool.engine.pipeline import run, collect_source_alerts
from essay_tool.engine.alerts import AlertCollector
from essay_tool.engine.doc_index import build_ops, body_list_field_kinds, body_toc_field_parts
from essay_tool.config.preset_io import load_preset
from essay_tool.gui import adjust_preview as ap

BASE = load_preset(ROOT / "presets" / "中国政法大学-本科毕业论文.json")
T = pathlib.Path(tempfile.mkdtemp())
ALL = (set(ph.PLACEHOLDER_PROMPTS.values()) | set(ph.CONTENT_PLACEHOLDER_PROMPTS.values())
       | set(ph.ENGLISH_PLACEHOLDER_PROMPTS.values()) | {ph.ANCHOR_CONTENT_PROMPTS[k] for k in ("参考文献", "附录", "致谢")}
       | {ph.APPENDIX_BODY_PROMPT})
NAME = {}
for tbl in (ph.PLACEHOLDER_PROMPTS, ph.CONTENT_PLACEHOLDER_PROMPTS, ph.ENGLISH_PLACEHOLDER_PROMPTS):
    for k, v in tbl.items(): NAME[v] = k
NAME[ph.ANCHOR_CONTENT_PROMPTS["参考文献"]] = "参考文献占位"; NAME[ph.ANCHOR_CONTENT_PROMPTS["致谢"]] = "致谢占位"
NAME[ph.ANCHOR_CONTENT_PROMPTS["附录"]] = "附录题名占位"; NAME[ph.APPENDIX_BODY_PROMPT] = "附录内容占位"

TEMPLATES = {
    "只有正文": ["一、绪论", "这是正文。", "二、分析", "更多正文。"],
    "有题目和摘要": ["论文题目", "摘要", "摘要正文。", "关键词", "甲；乙", "一、绪论", "这是正文。", "参考文献", "[1] 某某. 论文. 2020."],
    "只有标签没内容": ["论文题目", "摘要", "关键词", "一、绪论", "这是正文。", "参考文献", "致谢"],
}

def counts(path):
    c = collections.Counter()
    for p in Document(str(path)).paragraphs:
        t = p.text.strip()
        if t in ALL: c[NAME[t]] += 1
    return c

def export_auto(src, wd):
    return run(src, None, wd, alerts=AlertCollector(), preset_config=BASE, check_source_alerts=False)

def export_gui(src, wd):
    rows, index = ap.build_rows(src)
    ap.seed_enabled_from(BASE.structure_required, rows)
    for k in list(ap.ENABLED):
        if k != "插图和附表清单": ap.ENABLED[k] = True
    doc = Document(str(src))
    _c, plan, _cap = collect_source_alerts(document=doc, index=index, structure_required=BASE.structure_required,
        existing_list_kinds=body_list_field_kinds(doc), cover_end=None, regenerate_fields=frozenset(), include_source_checks=False)
    ops = build_ops(rows, index, enabled=dict(ap.ENABLED), insert_toc=("目录" not in body_toc_field_parts(doc)), list_kinds=plan)
    return run(src, None, wd, alerts=AlertCollector(), index=index, ops=ops, preset_config=BASE, check_source_alerts=False)

problems = []
n_cases = 0
for exporter_name, exporter in (("自动路径", export_auto), ("结构页路径", export_gui)):
    for tname, lines in TEMPLATES.items():
        for pos in range(len(lines) + 1):
            for n_blank in (1, 2):
                v = lines[:pos] + [""] * n_blank + lines[pos:]
                d = Document()
                for t in v: d.add_paragraph(t)
                src = T / f"s{n_cases}.docx"; d.save(str(src))
                n_cases += 1
                try:
                    o1 = exporter(src, T / f"w1_{n_cases}")
                    c1 = counts(o1)
                    # 回炉：用同一种路径再导一遍
                    o2 = exporter(pathlib.Path(o1), T / f"w2_{n_cases}")
                    c2 = counts(o2)
                except Exception as exc:
                    problems.append((exporter_name, tname, pos, n_blank, f"异常：{exc!r}")); continue
                dups = {k: n for k, n in c1.items() if n > 1}
                if dups:
                    problems.append((exporter_name, tname, pos, n_blank, f"第一遍就重复：{dups}"))
                elif c1 != c2:
                    diff = {k: (c1.get(k, 0), c2.get(k, 0)) for k in set(c1) | set(c2) if c1.get(k, 0) != c2.get(k, 0)}
                    problems.append((exporter_name, tname, pos, n_blank, f"再导一遍数量变了 {diff}"))
print(f"共扫描 {n_cases} 种情形（3 种文档 × 每个位置插 1/2 个空行 × 2 条导出路径）")
if not problems:
    print("全部通过：没有重复插入，回炉数量稳定")
else:
    print(f"发现 {len(problems)} 处问题：")
    by = collections.defaultdict(list)
    for p in problems: by[p[4]].append(p[:4])
    for msg, items in list(by.items())[:12]:
        print(f"  {msg}   × {len(items)}   例：{items[0]}")

sys.exit(1 if problems else 0)
