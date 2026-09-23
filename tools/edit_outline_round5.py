"""第 5 轮改表（用户 2026-09-16 逐条指定，全部用浅蓝色标出）。

按用户对上一轮报告 A/B/C 三节的批示执行：

- A1、A2：新增一章「7. 手动调整格式」，下辖 7.1 页面设置、7.2 格式设置（追加在表末，
  不移动任何现有行）。
- B7：把「状态栏」这个已经不存在的叫法统一为「日志区」（C4、C5、C15 三处）。
- C11：5.1 导出拆开写，导出论文与导出空白模板各成一段。
- C12：1.1 导入补上纯文本（.txt）的支持。
- C13：1.3 自动清理补上「条」「天/周/月」两种口径与「月按三十日计」。
- C15：3.3 补上结构调整不落盘。
- C16：6.3 崩溃补上错误代码的含义、崩溃日志位置与反馈方式。
- C14 与 D 节各项不动；6.2 正文不动。

另：新建 Sheet3，写入 6.2 的「其他告警」表（严肃、学术口径；凡提示内容含 {原因} 的，
在触发条件下列出几条常见原因供排查），供 C26 将来引用。

做法与前面几轮一致：片段替换先断言原文确实在、整格替换先断言原值、改过的格子填浅蓝、
写盘前备份。幂等：已经改过的片段找不到会报出来，但不会重复追加。
跑法：.venv\\Scripts\\python.exe tools\\edit_outline_round5.py
"""

from __future__ import annotations

import pathlib
import shutil

import openpyxl
from openpyxl.styles import Alignment, PatternFill

ROOT = pathlib.Path(__file__).resolve().parents[1]
TARGET = ROOT / "reference" / "工具操作逻辑写作表.xlsx"
BACKUP_DIR = ROOT / "work"
OUTLINE = "工具操作逻辑"
ALERT_SHEET = "Sheet3"
LIGHT_BLUE = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")

# ---------------------------------------------------------------- A1、A2 新增章

NEW_CHAPTER = ("7. 手动调整格式", "7.1 页面设置", "7.2 格式设置")

PAGE_SECTION = """页面页签用于设定纸张、页边距、页眉与页码，各项取值可在页签上直接修改，修改后于导出时生效。可设定的项目如下：
（1）纸张与方向：纸张可选 A4、A3、B5 或自定义；选择自定义时，另行填写宽与高，单位为厘米。方向可选纵向或横向。装订线以厘米计。
（2）页边距与距边：上、下、左、右页边距，以及页眉距边、页脚距边，均以厘米计。
（3）页眉：按文档划分出的各部分分别填写页眉文字；另有开关决定是否为除封面外的页眉应用装饰线。
（4）页脚与页码：按部分分别决定是否插入页码，并可指定页码格式、前缀与后缀；未勾选页码的部分，格式、前缀与后缀三项不生效。
导出时，文档按前置部分与正文部分分节，页眉与页码按部分分别生效，每一部分的页码均自 1 起算。"""

FORMAT_SECTION = """格式页签以表格形式列出论文中各角色（题目、各级标题、正文、摘要、图题、表格文字、脚注等）的字体、字号、段落格式等取值。表格按论文出现的先后分为前置部分、正文部分、正文以后、跨页元素四块，各块可折叠；第一行与第一列在滚动时固定。每一行对应一种角色，修改某一行即修改该类内容的排版。行不要求全部用到，只有实际用到的角色，才会在成品中生成对应的样式；名字之后带星号的行，其格式以直接格式写在段落上，不写入 Word 的样式面板。
表格之下另设引注区，放置表格未涵盖的两项：脚注的编号设置，以及尾注（文末参考文献列表）的条目编号。
在套用样式之前，工具会先清除原稿上手工设置的排版格式（字体、字号、颜色、行距、缩进、对齐等），以免其盖过样式。字符强调在清除时予以保留，默认为加粗、斜体、下划线、上下标，保留哪些由预设设定。
页签上修改的取值只记在本次运行之中，如需长期使用，须到文件页签点“保存当前设置为新配置”。"""

# ---------------------------------------------------------------- C 节各项（正文改动）

C21_OLD = ("选择导出论文后，工具将读取目前的各项预设和识别结果，导出论文并自动为其排版；"
           "导入空白模板的，工具将创建新的文档，应用各项预设和结构页的结构设置。")
C21_NEW = """点“导出论文到……”后，工具读取当前的各项预设与识别结果，生成成品并自动排版，随后写入你所选的位置。
点“导出空白模板到……”后，工具另行创建一份新文档，其中应用各项预设与结构页的结构设置，正文位置以占位文字标出，供你自行填写；封面是否插入，依封面页签的开关与所选形态而定。生成空白模板不要求先应用预设。"""

C27_OLD = ("如果工具出现异常而崩溃，将会弹出崩溃弹窗。\n"
           "只要你未曾删除原文件，你的数据安全不会受到影响。如果你已经删除了原文件但打开了备份，你可以前往查看备份。")
C27_NEW = """如果工具出现异常而崩溃，将会弹出崩溃弹窗，弹窗中给出错误代码与备份路径，并提供联系开发者的入口。错误代码由异常名称与一小段校验值构成，同一处故障每次出现时其代码相同，向开发者反馈时提供该代码即可。
崩溃的完整记录追加写入本用户目录下的 crash.log，即 %APPDATA%\\AutoEssayTypesettingTool\\crash.log。
只要你未曾删除原文件，你的数据安全不会受到影响。如果你已经删除了原文件但打开了备份，你可以前往查看备份。"""

# (单元格, 原文片段, 新片段) —— 片段在格内必须唯一
FRAGMENT_OPS = [
    # C12：1.1 导入补纯文本
    ("C2", "则会同时触发自动备份，见下文。",
     "则会同时触发自动备份，见下文。\n"
     "导入对话框默认只列 Word 文档（.docx）；将文件类型切换为“全部文件”，亦可选择纯文本（.txt），"
     "工具会逐行读入并新建一份文档，其后照常识别与排版。"),
    # B7：状态栏 -> 日志区
    ("C4", "并于状态栏提示错误原因。", "并于日志区提示错误原因。"),
    ("C5", "而只在状态栏告警。", "而只在日志区告警。"),
    ("C15", "关闭后也可以在左下侧状态栏再次查看。", "关闭后也可以在左下侧的日志区再次查看。"),
    # C13：1.3 自动清理补两种口径
    ("C4", "超过指定条数或日期的备份文件将被从旧到新删除。\n系统依靠内置清单",
     "超过指定条数或日期的备份文件将被从旧到新删除。\n"
     "单位为“条”时，系统按修改时间排序，仅保留最近的若干份，其余全部删除；"
     "单位为“天”“周”“月”时，系统删除修改时间早于当前时间减去该时长的全部备份，不考虑备份总份数。"
     "其中“月”以三十日计算，并非精确的日历月。\n系统依靠内置清单"),
    # C15：3.3 结构调整不落盘（追加在 3.3 末句之后）
    ("C11", "这意味着它们将会被列入插图清单或者附表清单（如有）。",
     "这意味着它们将会被列入插图清单或者附表清单（如有）。\n"
     "结构调整只保存在本次运行之中，不写入文档，也不写入任何设置文件；关闭程序后重新导入，须重新调整。"
     "更换文档时，前一份文档的调整一并作废。"),
]

# (单元格, 原值, 新值) —— 整格替换，先断言原值
CELL_OPS = [
    ("C21", C21_OLD, C21_NEW),      # C11：导出拆开写
    ("C27", C27_OLD, C27_NEW),      # C16：崩溃写细
]

# ---------------------------------------------------------------- 6.2 的表（新表）

ALERT_HEADERS = ("序号", "告警名称", "触发条件", "警告内容")
ALERT_ROWS: list[tuple[str, str, str, str]] = [
    ("1", "备份未完成",
     "导入完成后，系统对原文件同步执行自动备份。下列情形可能致其未能完成：目标磁盘剩余空间不足；"
     "备份目录不允许写入，例如设于系统盘或只读位置；备份目录已被移走或改名；原文件正被其他程序占用而无法读取。"
     "自动清理与备份共用同一条提示，故清理失败亦报此条",
     "备份原文件失败：{原因}（不影响导入）"),
    ("2", "备份设置未能保存",
     "备份开关、备份路径或自动清理选项写不进本用户的设置文件。下列情形可能致其未能保存："
     "用户目录不可写；设置文件正被其他程序占用；磁盘剩余空间不足",
     "备份设置保存失败（本次会话内仍生效）：{原因}"),
    ("3", "备份份数不合法",
     "在“自动删除过早备份”的份数输入框中填入 0、负数或非数字",
     "自动删除过早备份的数量须为正整数；{所填内容} 无法识别，已恢复为 {恢复后的数值}。"),
    ("4", "备份路径无法打开",
     "点“打开备份路径”时，备份目录无法建立，或无法在文件资源管理器中打开。下列情形可能致其失败："
     "路径所在磁盘已被移除；当前账户无权访问该路径；路径名含非法字符",
     "打开备份路径失败：{原因}"),
    ("5", "预设读取失败",
     "工具读取预设文件时失败。下列情形可能致其失败：文件已损坏或被手工改错；文件中缺少必需字段；"
     "文件正被其他程序占用；文件已被移动或删除",
     "读取预设“{预设名}”失败，结构部件启用状态按默认值显示：{原因}／（预设读取失败：{原因}）"),
    ("6", "预设应用失败",
     "点“应用”时该预设无法读取，其原因与前条相同",
     "配置应用失败：{预设名}，参数读取失败（{原因}）"),
    ("7", "预设保存失败",
     "点“保存当前设置为新配置”时，当前表单的取值不符要求，或预设文件无法写入。下列情形可能致其失败："
     "名称为空或含非法字符；某项取值超出允许范围；预设目录无权写入",
     "保存失败：设置内容有误（{原因}）／保存失败：{原因}"),
    ("8", "预设导入失败",
     "从 Excel 导入配置文件时读取失败。下列情形可能致其失败：工作表名称或表头与工具读取格式不符；"
     "单元格取值无法解析；文件不是 .xlsx，或文件已损坏",
     "导入失败：{原因}"),
    ("9", "预设删除失败",
     "删除自建预设时文件无法移除。下列情形可能致其失败：文件正被其他程序占用；目录无权写入。"
     "内置预设不可删除，系统不执行删除，另行提示",
     "删除失败：{原因}／内置配置不可删除。"),
    ("10", "预设导出失败",
     "把预设导出为 Excel 时，读取该预设或写入目标文件失败。下列情形可能致其失败：预设文件已损坏；"
     "目标文件正被 Excel 打开；目标位置无权写入",
     "导出失败：读取 {预设名} 时出错（{原因}）／导出失败：{原因}"),
    ("11", "空白模板生成失败",
     "点“导出空白模板到……”后，生成过程失败。下列情形可能致其失败：目标文件正被 Word 打开；"
     "目标位置无权写入；磁盘剩余空间不足",
     "生成空白模板失败：{原因}"),
    ("12", "论文生成失败",
     "点“导出论文到……”后，生成过程失败。下列情形可能致其失败：原稿或目标文件正被 Word 打开；"
     "目标位置无权写入；磁盘剩余空间不足；结构页的调整结果无法换算，此项并入本条",
     "生成论文失败：{原因}"),
    ("13", "结构识别失败",
     "导入文档后切至结构页签，工具自动识别文档时失败。下列情形可能致其失败：文档仍被 Word 打开；"
     "文件已损坏或并非有效的 .docx；文档设置了打开密码",
     "结构识别失败：{原因}／识别失败：{原因}"),
    ("14", "操作前提未满足",
     "未导入论文即点导出；封面总开关关闭时点封面内容区；点内置预设的删除按钮",
     "请先导入论文再导出！／封面总开关未开启，请先勾选为论文插入封面。／内置配置不可删除。"),
    ("15", "覆盖原文件前的确认",
     "导出位置选中的就是原文件本身，系统先行询问；取消则该次导出终止并记为下句",
     "弹窗询问是否覆盖，取消后提示：已取消导出（没有覆盖原文件）。"),
    ("16", "规范原文相关操作失败",
     "点“查阅《规范》原文”时窗口无法打开；点“另存官方版”时原始文件不存在，或另存失败",
     "规范浏览窗口打开失败：{原因}／规范文件另存失败：reference 目录中未找到原始文件。／另存失败：{原因}"),
    ("17", "程序崩溃",
     "出现未预料的异常，程序无法继续运行",
     "弹窗“哎呀！出现了一点小问题，程序已崩溃。”，并给出错误代码与备份目录；"
     "崩溃详情写入 %APPDATA%\\AutoEssayTypesettingTool\\crash.log。详见 6.3"),
]


def next_backup() -> pathlib.Path:
    for i in range(1, 100):
        candidate = BACKUP_DIR / f"_logic_sheet_backup_before_round5-{i}.xlsx"
        if not candidate.exists():
            return candidate
    raise SystemExit("备份文件名用尽")


def apply_fragment_ops(ws, notes: list[str], troubles: list[str]) -> list[str]:
    changed: list[str] = []
    for coord, old, new in FRAGMENT_OPS:
        cell = ws[coord]
        text = "" if cell.value is None else str(cell.value)
        if new in text:
            troubles.append(f"{OUTLINE}!{coord} 已经改过（跳过）：{old[:18]}…")
            continue
        if text.count(old) != 1:
            troubles.append(f"{OUTLINE}!{coord} 里找不到唯一片段（出现 {text.count(old)} 次）：{old[:24]}…")
            continue
        cell.value = text.replace(old, new, 1)
        cell.fill = LIGHT_BLUE
        changed.append(f"{OUTLINE}!{coord}（片段）{old[:14]}… -> 已改")
    notes.append(f"片段替换 {len(changed)} 处")
    return changed


def apply_cell_ops(ws, notes: list[str], troubles: list[str], changed: list[str]) -> None:
    for coord, old, new in CELL_OPS:
        cell = ws[coord]
        text = "" if cell.value is None else str(cell.value)
        if text == new:
            troubles.append(f"{OUTLINE}!{coord} 已经改过（跳过）")
            continue
        if text != old:
            troubles.append(f"{OUTLINE}!{coord} 的原值与预期不符，未改动：{text[:30]}…")
            continue
        cell.value = new
        cell.fill = LIGHT_BLUE
        changed.append(f"{OUTLINE}!{coord}（整格）已按新稿改写")


def add_chapter(ws, notes: list[str], troubles: list[str], changed: list[str]) -> None:
    title, sec1, sec2 = NEW_CHAPTER
    existing = [cell.value for cell in ws["A"]
                if cell.value is not None and str(cell.value).strip() == title]
    if existing:
        troubles.append(f"{OUTLINE}!A 列里已经有「{title}」，不重复新增")
        return
    row1 = ws.max_row + 1
    occupied = [(r, c, ws.cell(row=r, column=c).value)
                for r in (row1, row1 + 1) for c in (1, 2, 3)]
    if any(value is not None for _r, _c, value in occupied):
        troubles.append(f"{OUTLINE}!A{row1} 起的目标位置不是空的，未新增章节：{occupied}")
        return
    ws.cell(row=row1, column=1, value=title)
    ws.cell(row=row1, column=2, value=sec1)
    ws.cell(row=row1, column=3, value=PAGE_SECTION)
    ws.cell(row=row1 + 1, column=2, value=sec2)
    ws.cell(row=row1 + 1, column=3, value=FORMAT_SECTION)
    for row in (row1, row1 + 1):
        for col in (1, 2, 3):
            ws.cell(row=row, column=col).fill = LIGHT_BLUE
    ws.merge_cells(start_row=row1, start_column=1, end_row=row1 + 1, end_column=1)
    changed.append(f"{OUTLINE}!A{row1}:A{row1 + 1} 新增「{title}」，下辖 {sec1}、{sec2}")


def build_alert_sheet(book, notes: list[str], changed: list[str]) -> None:
    if ALERT_SHEET in book.sheetnames:
        del book[ALERT_SHEET]
    ws = book.create_sheet(ALERT_SHEET)
    for col, head in enumerate(ALERT_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=head)
        cell.fill = LIGHT_BLUE
    for i, row in enumerate(ALERT_ROWS, start=2):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.fill = LIGHT_BLUE
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABCD", (6, 20, 62, 46)):
        ws.column_dimensions[col].width = width
    changed.append(f"{ALERT_SHEET} 新表：{ALERT_HEADERS} + {len(ALERT_ROWS)} 条提示")
    notes.append(f"{ALERT_SHEET} 共 {len(ALERT_ROWS)} 行")


def main() -> None:
    book = openpyxl.load_workbook(TARGET)
    ws = book[OUTLINE]
    notes: list[str] = []
    troubles: list[str] = []
    changed: list[str] = []

    changed += apply_fragment_ops(ws, notes, troubles)
    apply_cell_ops(ws, notes, troubles, changed)
    add_chapter(ws, notes, troubles, changed)
    build_alert_sheet(book, notes, changed)

    if troubles:
        print("⚠ 需要注意：")
        for item in troubles:
            print(f"   {item}")
        print()

    if not changed:
        print("没有任何改动。")
        return

    backup = next_backup()
    shutil.copy2(TARGET, backup)
    print(f"改前版本已备份：{backup.relative_to(ROOT)}\n")
    print(f"共 {len(changed)} 项改动（均填浅蓝）：")
    for item in changed:
        print(f"  {item}")

    book.save(TARGET)
    print(f"\n已写回 {TARGET.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
