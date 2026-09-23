"""格式页「清理规则」块与「标签写法」里「另起一页」复选框的界面检查（隐藏窗口，不弹界面）：
默认状态、勾选后收集进预设、保存再载入还原、跟随样式联动、从界面导出端到端生效。

跑法：.venv\\Scripts\\python.exe tools\\check_gui_rules_20260920.py    （退出码 0＝全部通过）
"""

import sys, tempfile, pathlib, tkinter as tk
sys.stdout.reconfigure(encoding="utf-8")
ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
orig_init, orig_state = tk.Tk.__init__, tk.Tk.state
tk.Tk.__init__ = lambda self, *a, **k: (orig_init(self, *a, **k), self.withdraw())[0]
tk.Tk.state = lambda self, *a: "withdrawn"
from tkinter import filedialog, messagebox, ttk
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
from essay_tool.gui import main_window as mw
from essay_tool.config.preset_io import load_preset, save_preset_json

bad = 0
def check(n, c, extra=""):
    global bad
    print(("OK   " if c else "FAIL ") + n, "" if c else extra); bad += (not c)

T = pathlib.Path(tempfile.mkdtemp())
app = mw.MainWindow()
def settle():
    for _ in range(3): app.update_idletasks(); app.update()
def walk(w):
    yield w
    for c in w.winfo_children(): yield from walk(c)
try:
    messagebox.askokcancel = messagebox.askyesno = lambda *a, **k: True
    messagebox.showinfo = messagebox.showwarning = messagebox.showerror = lambda *a, **k: None
    settle()
    # ---- 1 默认状态
    check("清理规则块存在：分页符/空行默认不勾", not app.v_clear_pb.get() and not app.v_clear_blank.get())
    check("保留字符格式默认：加粗/斜体/下划线/上下标勾选，颜色不勾",
          {k: v.get() for k, v in app.keep_vars.items()} == {"b": True, "i": True, "u": True, "vertAlign": True, "color": False},
          {k: v.get() for k, v in app.keep_vars.items()})
    cfg0 = app._collect_current_preset_config()
    check("默认收集：清除开关为假、无标签覆盖", (cfg0.clear_page_breaks, cfg0.clear_blank_lines, cfg0.label_page_break) == (False, False, {}), (cfg0.clear_page_breaks, cfg0.clear_blank_lines, cfg0.label_page_break))
    check("默认收集：保留格式＝默认四项", cfg0.keep_run_formatting == frozenset({"b", "i", "u", "vertAlign"}), cfg0.keep_run_formatting)
    # 标签复选框默认跟样式：参考文献套一级标题样式（本来分页），摘要标签本来不分页
    check("标签复选框跟随样式：参考文献=勾、摘要=不勾", app.label_pb_vars["参考文献"].get() is True and app.label_pb_vars["摘要"].get() is False,
          (app.label_pb_vars["参考文献"].get(), app.label_pb_vars["摘要"].get()))
    check("标签复选框覆盖了全部标签部件（含引言/结论）", set(app.label_pb_vars) == set(mw.LABEL_PART_NAMES) | set(mw.HEADING_LABEL_PARTS), sorted(app.label_pb_vars))
    # ---- 2 改动 → 收集
    def click(cb):
        cb.invoke(); settle()
    cbs = [w for w in walk(app) if isinstance(w, ttk.Checkbutton)]
    def find_cb(var):
        return next(c for c in cbs if str(c.cget("variable")) == str(var))
    click(find_cb(app.v_clear_pb)); click(find_cb(app.v_clear_blank))
    click(find_cb(app.keep_vars["color"]))
    click(find_cb(app.keep_vars["i"]))                          # 去掉斜体
    click(find_cb(app.label_pb_vars["摘要"]))                   # 摘要 → 另起一页（与样式不同）
    click(find_cb(app.label_pb_vars["参考文献"]))               # 参考文献 → 不另起（与样式不同）
    click(find_cb(app.label_pb_vars["致谢"])); click(find_cb(app.label_pb_vars["致谢"]))   # 点两下回到样式原值 → 不该记覆盖
    cfg = app._collect_current_preset_config()
    check("收集：清除分页符/空行", (cfg.clear_page_breaks, cfg.clear_blank_lines) == (True, True))
    check("收集：保留颜色、不保留斜体", cfg.keep_run_formatting == frozenset({"b", "u", "vertAlign", "color"}), cfg.keep_run_formatting)
    check("收集：只记与样式不同的覆盖（致谢点两下回原值不记）", cfg.label_page_break == {"摘要": True, "参考文献": False}, cfg.label_page_break)
    # ---- 3 保存 → 换一份预设 → 再载入：界面还原
    j = T / "mine.json"; save_preset_json(cfg, j)
    app._load_formats_from(load_preset(ROOT / "presets" / "中国政法大学-课程考核.json"), "课程考核"); settle()
    check("换成别的预设后：开关回到该预设的值（关）", not app.v_clear_pb.get() and not app.v_clear_blank.get())
    app._load_formats_from(load_preset(j), "我的"); settle()
    check("载入保存的预设：开关还原", app.v_clear_pb.get() and app.v_clear_blank.get())
    check("载入保存的预设：保留格式还原（含颜色）", app._collect_keep_run_formatting(None) == frozenset({"b", "u", "vertAlign", "color"}))
    check("载入保存的预设：标签覆盖还原", app.label_pb_vars["摘要"].get() is True and app.label_pb_vars["参考文献"].get() is False)
    check("再收集一遍结果不变（往返稳定）", app._collect_current_preset_config().label_page_break == {"摘要": True, "参考文献": False})
    # ---- 4 跟随样式：改样式表"段前分页"，没动过的标签复选框跟着变，动过的不变
    s02 = next(rv for rv in app.fmt_rows if rv["id"] == "S02")
    before_zh = app.label_pb_vars["引言"].get()
    s02["pb"].set(not s02["pb"].get()); settle()
    check("改一级标题样式的段前分页：没动过的「引言」标签跟着变", app.label_pb_vars["引言"].get() == (not before_zh), (before_zh, app.label_pb_vars["引言"].get()))
    check("……动过的「参考文献」不受影响（仍是预设里写的不分页）", app.label_pb_vars["参考文献"].get() is False)
    s02["pb"].set(not s02["pb"].get()); settle()
    # ---- 5 端到端：从界面点导出，清理规则真的生效
    W = nsdecls("w")
    d = Document()
    for t in ("一、绪论", "正文一。", "", "", "正文二。"): d.add_paragraph(t)
    pb = d.add_paragraph(""); pb._p.append(parse_xml(f'<w:r {W}><w:br w:type="page"/></w:r>'))
    d.add_paragraph("二、分析"); d.add_paragraph("正文三。")
    src = T / "in.docx"; d.save(str(src))
    app.var_agree.set(True)
    filedialog.askopenfilename = lambda *a, **k: str(src)
    app.on_import_file(); settle()
    out = T / "out.docx"
    filedialog.asksaveasfilename = lambda *a, **k: str(out)
    app.on_export_doc(); settle()
    o = Document(str(out))
    n_br = sum(1 for p in o.paragraphs for b in p._p.iter(qn("w:br")) if b.get(qn("w:type")) == "page")
    empties = [p for p in o.paragraphs[:8] if not p.text]
    check("端到端：导出成功", out.exists())
    check("端到端：成品里没有手打分页符", n_br == 0, n_br)
    body_texts = [p.text for p in o.paragraphs]
    i = body_texts.index("正文一。")
    check("端到端：正文一、正文二之间的两个空行被清掉", body_texts[i + 1] == "正文二。", body_texts[i:i + 3])
    check("端到端：分页符所在的空段也被清掉（正文二后面直接是二、分析）", body_texts[body_texts.index("正文二。") + 1].startswith("二、"), body_texts[i:i + 5])
finally:
    app.destroy()
    tk.Tk.__init__, tk.Tk.state = orig_init, orig_state
print("失败:", bad)

sys.exit(1 if bad else 0)
