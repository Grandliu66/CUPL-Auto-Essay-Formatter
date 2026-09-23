"""2026-09-20 三处修复的回归检查（用户报的 bug）——改动这几处代码后跑一遍：

    .venv\\Scripts\\python.exe tools\\check_fixes_20260920.py

1. **占位语不许插两遍**：结构页（面板）路径下，摘要/关键词/英文摘要/英文关键词/参考文献/
   附录/致谢的"标签位"与"内容位"两条生成指令，内容占位语只能各出现一次，且标签紧跟内容；
   把成品再导一遍也不增长。
2. **参考文献里的链接保持原有颜色/划线**：真实文件 `reference/4√正文9.1.docx` 里的域式
   超链接（rStyle=16 + color=000000 + u=none）导出后不能变蓝；合成的几种链接写法
   （w:hyperlink、域式、手染网址）也要保住颜色与划线，字号等其余直接格式照常清掉。
3. **导出按钮的可用状态只跟「我已阅读并同意」走**：导入→勾选→导出→取消导入→再导入→再导出，
   全程不需要重新勾选（原来取消导入会把按钮置灰而勾选框还勾着）。

退出码 0＝全部通过。GUI 部分用隐藏窗口，不会弹出任何界面。
"""

from __future__ import annotations

import collections
import pathlib
import sys
import tempfile
import tkinter as tk

sys.stdout.reconfigure(encoding="utf-8")
ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document                                           # noqa: E402
from docx.oxml import parse_xml                                     # noqa: E402
from docx.oxml.ns import nsdecls, qn                                # noqa: E402

from essay_tool.config.preset_io import load_preset                 # noqa: E402
from essay_tool.engine import placeholders as ph                    # noqa: E402
from essay_tool.engine.alerts import AlertCollector                 # noqa: E402
from essay_tool.engine.doc_index import (body_list_field_kinds,     # noqa: E402
                                         body_toc_field_parts, build_ops)
from essay_tool.engine.pipeline import collect_source_alerts, run   # noqa: E402
from essay_tool.gui import adjust_preview as ap                     # noqa: E402

bad = 0


def check(name: str, ok: bool, extra="") -> None:
    global bad
    print(("OK   " if ok else "FAIL ") + name, "" if ok else extra)
    bad += (not ok)


T = pathlib.Path(tempfile.mkdtemp())
PRESET = load_preset(ROOT / "presets" / "中国政法大学-本科毕业论文.json")


# ============================================================ 1 占位语不重复
def gui_export(src: pathlib.Path):
    """走结构页的路径：build_rows → ENABLED → build_ops → run(index, ops)。"""
    rows, index = ap.build_rows(src)
    ap.seed_enabled_from(PRESET.structure_required, rows)
    for k in list(ap.ENABLED):
        if k != "插图和附表清单":
            ap.ENABLED[k] = True
    doc = Document(str(src))
    _c, list_plan, _cap = collect_source_alerts(
        document=doc, index=index, structure_required=PRESET.structure_required,
        existing_list_kinds=body_list_field_kinds(doc), cover_end=None,
        regenerate_fields=frozenset(), include_source_checks=False)
    ops = build_ops(rows, index, enabled=dict(ap.ENABLED),
                    insert_toc=("目录" not in body_toc_field_parts(doc)), list_kinds=list_plan)
    wd = T / "w"
    wd.mkdir(exist_ok=True)
    return pathlib.Path(run(src, None, wd, alerts=AlertCollector(), index=index, ops=ops,
                            preset_config=PRESET, check_source_alerts=False))


def make_src(kind: str) -> pathlib.Path:
    d = Document()
    lines = {"body_only": ("一、绪论", "这是正文。", "二、分析", "更多正文。"),
             "labels_no_body": ("摘要", "关键词", "一、绪论", "这是正文。", "参考文献", "致谢")}[kind]
    for t in lines:
        d.add_paragraph(t)
    p = T / f"{kind}.docx"
    d.save(str(p))
    return p


ALL_PLACEHOLDERS = (set(ph.PLACEHOLDER_PROMPTS.values()) | set(ph.CONTENT_PLACEHOLDER_PROMPTS.values())
                    | set(ph.ENGLISH_PLACEHOLDER_PROMPTS.values())
                    | {ph.ANCHOR_CONTENT_PROMPTS[k] for k in ("参考文献", "附录", "致谢")}
                    | {ph.APPENDIX_BODY_PROMPT})
LABEL_THEN_CONTENT = {   # 标签文字 → 紧跟其后的占位语（成品里必须一前一后）
    "【摘要】": ph.CONTENT_PLACEHOLDER_PROMPTS["摘要"],
    "【关键词】": ph.CONTENT_PLACEHOLDER_PROMPTS["关键词"],
    "【ABSTRACT】": ph.ENGLISH_PLACEHOLDER_PROMPTS["英文摘要"],
    "【KEYWORDS】": ph.ENGLISH_PLACEHOLDER_PROMPTS["英文关键词"],
    "致  谢": ph.ANCHOR_CONTENT_PROMPTS["致谢"],
}

for kind in ("body_only", "labels_no_body"):
    out = gui_export(make_src(kind))
    texts = [p.text.strip() for p in Document(str(out)).paragraphs]
    cnt = collections.Counter(t for t in texts if t in ALL_PLACEHOLDERS)
    dups = {t[:14]: n for t, n in cnt.items() if n > 1}
    check(f"[{kind}] 占位语没有重复出现", not dups, dups)
    for label, content in LABEL_THEN_CONTENT.items():
        if label in texts:
            i = texts.index(label)
            check(f"[{kind}] 「{label}」出现一次且紧跟内容占位", texts.count(label) == 1
                  and i + 1 < len(texts) and texts[i + 1] == content, texts[i:i + 2])
    # 成品再导一遍：不增长
    again = gui_export(out)
    t2 = [p.text.strip() for p in Document(str(again)).paragraphs]
    check(f"[{kind}] 成品再导一遍，占位语数量不变", collections.Counter(
        t for t in t2 if t in ALL_PLACEHOLDERS) == cnt,
        (sum(cnt.values()), sum(1 for t in t2 if t in ALL_PLACEHOLDERS)))

# ============================================================ 2 参考文献链接格式
W = nsdecls("w")


def colored(text, style=False, direct=True):
    rpr = ('<w:rStyle w:val="Hyperlink"/>' if style else "") + (
        '<w:color w:val="0563C1"/><w:sz w:val="26"/><w:u w:val="single"/>' if direct else "")
    return f'<w:r {W}><w:rPr>{rpr}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'


d = Document()
for t in ("一、绪论", "正文。", "参考文献"):
    d.add_paragraph(t)
p = d.add_paragraph("[1] 甲. 论A. ")
p._p.append(parse_xml(f'<w:hyperlink {W} w:anchor="x">{colored("https://example.com/A")}</w:hyperlink>'))
p = d.add_paragraph("[2] 乙. 论C. ")
for x in (f'<w:r {W}><w:fldChar w:fldCharType="begin"/></w:r>',
          f'<w:r {W}><w:instrText xml:space="preserve"> HYPERLINK "https://example.com/C" </w:instrText></w:r>',
          f'<w:r {W}><w:fldChar w:fldCharType="separate"/></w:r>', colored("https://example.com/C"),
          f'<w:r {W}><w:fldChar w:fldCharType="end"/></w:r>'):
    p._p.append(parse_xml(x))
p = d.add_paragraph("[3] 丙. 论D. ")
p._p.append(parse_xml(colored("https://example.com/D")))
src = T / "links.docx"
d.save(str(src))
out = Document(str(run(src, None, T / "w2", alerts=AlertCollector(), preset_config=PRESET,
                       check_source_alerts=False)))


def link_fmt(doc, needle):
    for r in doc.element.body.iter(qn("w:r")):
        if needle in "".join(t.text or "" for t in r.findall(qn("w:t"))):
            rpr = r.find(qn("w:rPr"))
            g = lambda tag: (rpr.find(qn(tag)).get(qn("w:val")) if rpr is not None
                             and rpr.find(qn(tag)) is not None else None)
            return g("w:color"), g("w:u"), g("w:sz")
    return None


for needle, label in (("example.com/A", "w:hyperlink 里的链接"), ("example.com/C", "域式超链接"),
                      ("example.com/D", "手动染色的网址")):
    fmt = link_fmt(out, needle)
    check(f"{label}：保留颜色与划线，字号照常清掉", fmt == ("0563C1", "single", None), fmt)

real = ROOT / "reference" / "4√正文9.1.docx"
if real.exists():
    def real_links(doc):
        res = []
        for pp in doc.paragraphs:
            if not any("HYPERLINK" in (t.text or "") for t in pp._p.iter(qn("w:instrText"))):
                continue
            in_res = False
            for r in pp._p.iter(qn("w:r")):
                fc = r.find(qn("w:fldChar"))
                if fc is not None:
                    k = fc.get(qn("w:fldCharType"))
                    in_res = (k == "separate") if k in ("separate", "end") else in_res
                    continue
                t = "".join(x.text or "" for x in r.findall(qn("w:t")))
                if t and in_res:
                    rpr = r.find(qn("w:rPr"))
                    g = lambda tag: (rpr.find(qn(tag)).get(qn("w:val")) if rpr is not None
                                     and rpr.find(qn(tag)) is not None else None)
                    res.append((t, g("w:color"), g("w:u"), g("w:rStyle")))
        return res
    before = real_links(Document(str(real)))
    after = real_links(Document(str(run(real, None, T / "w3", alerts=AlertCollector(),
                                        preset_config=PRESET, check_source_alerts=False))))
    check(f"真实文件 4√正文9.1：{len(before)} 处超链接的颜色/划线/字符样式原样保留（不变蓝）",
          bool(before) and before == after, (before[:2], after[:2]))
else:
    print("SKIP 真实文件 reference/4√正文9.1.docx 不在")

# ============================================================ 3 导出按钮状态
orig_init, orig_state = tk.Tk.__init__, tk.Tk.state
tk.Tk.__init__ = lambda self, *a, **k: (orig_init(self, *a, **k), self.withdraw())[0]
tk.Tk.state = lambda self, *a: "withdrawn"
from tkinter import filedialog, messagebox                          # noqa: E402
from essay_tool.gui import main_window as mw                        # noqa: E402

app = mw.MainWindow()
try:
    messagebox.askokcancel = messagebox.askyesno = lambda *a, **k: True
    messagebox.showinfo = messagebox.showwarning = messagebox.showerror = lambda *a, **k: None
    st = lambda: (str(app.btn_exp_doc.cget("state")), str(app.btn_exp_blank.cget("state")))
    SRC = ROOT / "reference" / "民总.docx"
    check("未勾选协议 → 导出按钮是灰的", st() == ("disabled", "disabled"), st())
    app.var_agree.set(True)
    app.update()
    check("勾选协议 → 导出按钮可用", st() == ("normal", "normal"), st())
    filedialog.askopenfilename = lambda *a, **k: str(SRC)
    filedialog.asksaveasfilename = lambda *a, **k: str(T / "o1.docx")
    app.on_import_file()
    app.on_export_doc()
    check("导入→导出后按钮仍可用", (T / "o1.docx").exists() and st() == ("normal", "normal"), st())
    app.on_cancel_import()
    check("取消导入后：仍是勾选状态，按钮仍可用（不需要重新勾选协议）",
          app.var_agree.get() and st() == ("normal", "normal"), (app.var_agree.get(), st()))
    app.on_import_file()
    filedialog.asksaveasfilename = lambda *a, **k: str(T / "o2.docx")
    app.on_export_doc()
    check("取消导入→重新导入→再导出：直接成功", (T / "o2.docx").exists())
    app.var_agree.set(False)
    app.update()
    check("取消勾选 → 按钮变灰（协议闸门仍有效）", st() == ("disabled", "disabled"), st())
finally:
    app.destroy()
    tk.Tk.__init__, tk.Tk.state = orig_init, orig_state

print("失败:", bad)
sys.exit(1 if bad else 0)
