"""用法：.venv\\Scripts\\python.exe tools\\inspect_docx.py <docx路径>

打印段落样式、页面设置，用于人工核对排版结果。
"""

import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from docx import Document
from docx.oxml.ns import qn


def main():
    path = pathlib.Path(sys.argv[1])
    doc = Document(path)

    s = doc.sections[0]
    print(f"页面: {s.page_width.cm:.2f}x{s.page_height.cm:.2f}cm 方向={s.orientation} "
          f"边距(上下左右)={s.top_margin.cm:.2f}/{s.bottom_margin.cm:.2f}/{s.left_margin.cm:.2f}/{s.right_margin.cm:.2f}cm")

    normal = doc.styles["Normal"]
    print(f"Normal 样式: 字号={normal.font.size.pt}pt 加粗={normal.font.bold} 对齐={normal.paragraph_format.alignment}")
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    print(f"  字体: ascii={rFonts.get(qn('w:ascii'))} eastAsia={rFonts.get(qn('w:eastAsia'))}")
    pPr = normal.element.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    print(f"  缩进属性: {dict(ind.attrib) if ind is not None else None}")
    spacing = pPr.find(qn("w:spacing"))
    print(f"  行距属性: {dict(spacing.attrib) if spacing is not None else None}")

    if "--styles" in sys.argv:
        print()
        print("样式清单（styleId / 名称 / 是否是内置）：")
        for style in doc.styles:
            try:
                builtin = style.builtin
            except Exception:
                builtin = "?"
            print(f"  {style.style_id!r:<24} {style.name!r:<16} 内置={builtin}")

    print()
    for i, p in enumerate(doc.paragraphs):
        print(f"[{i}] 样式={p.style.name!r} 文字={p.text[:30]!r}")

    # 表格单元格里的段落不在 doc.paragraphs 里，要单独走一遍——不然套格式出问题
    # 时完全看不出来。
    if doc.tables:
        print()
        print(f"表格：共 {len(doc.tables)} 个")
        for ti, table in enumerate(doc.tables):
            # 首行有没有设"标题行重复"（续表要求）——这个只能直接看 XML，
            # python-docx 没有对应属性。
            repeat = False
            if table.rows:
                trPr = table.rows[0]._tr.find(qn("w:trPr"))
                repeat = trPr is not None and trPr.find(qn("w:tblHeader")) is not None
            print(f"  表{ti}: 首行标题行重复={repeat}")
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    for p in cell.paragraphs:
                        print(f"[表{ti}-{ri}-{ci}] 样式={p.style.name!r} 文字={p.text[:20]!r}")


if __name__ == "__main__":
    main()
