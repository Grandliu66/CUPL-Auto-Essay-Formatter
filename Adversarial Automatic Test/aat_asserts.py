# -*- coding: utf-8 -*-
r"""断言式覆盖：不走识别索引、**直接查产物**的那几条规则。

路径 a（喂期望识别结果）够不着的，用"直接断言"补（README「下一步」第 3 项）：

  * **S4 附录序号题名**：`附录A …` 这类段落要套 S20（附录序号题名）样式。
    识别在 `detect.detect_appendix_entries_from_index()`（范围限制很死：
    "附录"标题之后、下一个锚点之前；"附录+一个拉丁字母"开头、不以句末标点
    结尾、≤40 字；字间与词尾空格允许）。它不经过 `Entry.part`——所以路径 a
    看不见它，只能断言成品里的样式。
  * **清单域插入计划**：勾「插图清单」→ 成品里应出现按样式名收题注的域指令
    `TOC \h \z \t "图题,1"`（`sections._LIST_FIELD_CODE`）；源文档里没有
    对应题注时**不该**出现该域（空清单没意义）。
  * **标签前"另起一页"覆盖**（逻辑表总则四，2026-09-20）：预设
    `label_page_break` 写了某个部件就以它为准（True／False 两个方向都要压过
    样式的"段前分页"）；没写就跟随样式。缺失补件插出来的标签同样生效。
  * **清理规则：手打分页符／空行**（总则三，2026-09-20）：封面范围以外清；
    封面范围里连空段带分页符都保住；两张表格之间、文末表格之后各留的空段
    不删；只含手打分页符的段落，仅在同时勾选"清除分页符"时才按空行清掉。
  * **保留的字符格式**（总则一，2026-09-20 增"颜色"）：默认保留加粗／斜体／
    下划线／上下标、颜色清掉；勾上颜色则颜色也留；"全清"一个不留；一至四级
    标题连加粗／颜色都不留；参考文献里的链接文字始终保留颜色与划线。

合成夹具在本脚本里现造（不依赖某份具体 reference 稿）。

用法：.venv\Scripts\python.exe "Adversarial Automatic Test\aat_asserts.py"
"""
from __future__ import annotations

import pathlib
import shutil
import sys
import time

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent
for entry in (str(ROOT), str(HERE)):
    if entry not in sys.path:
        sys.path.insert(0, entry)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document                                       # noqa: E402
from docx.enum.text import WD_BREAK                             # noqa: E402
from docx.oxml.ns import qn                                     # noqa: E402
from docx.shared import RGBColor                                # noqa: E402
from essay_tool.config.preset_io import load_preset             # noqa: E402
from essay_tool.engine import pipeline                          # noqa: E402
from essay_tool.engine.alerts import AlertCollector             # noqa: E402

PRESET = ROOT / "presets" / "中国政法大学-本科毕业论文.json"
WORK = HERE / "_out" / "asserts"
REPORT = HERE / "_out" / "aat-asserts.txt"

fails: list[str] = []


def check(name, ok, detail="", rule="") -> None:
    tail = f"（{rule}）" if rule else ""
    print(f"   {'OK  ' if ok else 'FAIL'}  {name}{tail}"
          + (f"  → {detail}" if detail else ""))
    if not ok:
        fails.append(f"{name}{tail}  {detail}")


def _write(path: pathlib.Path, lines: list[tuple[str, str | None]]) -> None:
    """`lines`：[(文本, 样式名或 None)]。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for text, style in lines:
        if style:
            document.add_paragraph(text, style=style)
        else:
            document.add_paragraph(text)
    document.save(str(path))


def _run(source: pathlib.Path, tag: str, *, enable: dict | None = None,
         tweak=None, cover_end: int | None = None):
    """`tweak(config)`：在跑之前改预设（清理开关、保留格式、标签页、样式值…）。
    `cover_end`：封面区间（排除识别／格式化的范围，清理规则要避开它）。"""
    config = load_preset(PRESET)
    if tweak is not None:
        tweak(config)
    if enable:
        for key, value in enable.items():
            config.structure_required[key] = value
    out = WORK / f"out-{tag}"
    shutil.rmtree(out, ignore_errors=True)
    product = pipeline.run(source, None, out, preset_config=config,
                           alerts=AlertCollector(), cover_end=cover_end)
    return product


def _instr_codes(path: pathlib.Path) -> list[str]:
    """产物里所有域指令：`w:fldSimple/@w:instr` 属性 + `w:instrText` 元素。

    ⚠️ 我们自己的清单域是 `fldSimple`——指令写在**属性**上，只扫 `instrText`
    会漏（2026-09-19 实测：域明明在，断言报"没插"）。
    """
    document = Document(str(path))
    codes: list[str] = []
    for element in document.element.body.iter():
        if element.tag.endswith("}fldSimple"):
            instr = element.get(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr")
            if instr:
                codes.append(instr)
    for field in document.element.body.iter():
        if field.tag.endswith("}instrText"):
            if field.text:
                codes.append(field.text)
    return codes


def _s20_mismatches(paragraph, spec) -> list[str]:
    """段落的**直接格式**跟预设 S20 的差多少（S20 是"直接格式"角色：不建独立
    样式——《styles._DIRECT_FORMAT_BASE》里 `"S20": "S01"`，格式直写段上，
    所以断言落在段落格式上，不看样式名——2026-09-19 实测踩过）。"""
    bad: list[str] = []
    pf = paragraph.paragraph_format
    if pf.alignment != spec.alignment:
        bad.append(f"对齐 {pf.alignment} ≠ {spec.alignment}")
    if paragraph.runs:
        run = paragraph.runs[0]
        if run.font.size != spec.size:
            bad.append(f"字号 {run.font.size} ≠ {spec.size}")
        if bool(run.font.bold) != bool(spec.bold):
            bad.append(f"加粗 {run.font.bold} ≠ {spec.bold}")
        rPr = run._element.find(qn("w:rPr"))
        rFonts = rPr.find(qn("w:rFonts")) if rPr is not None else None
        east = rFonts.get(qn("w:eastAsia")) if rFonts is not None else None
        if east != spec.font_east_asian:
            bad.append(f"中文字体 {east!r} ≠ {spec.font_east_asian!r}")
    else:
        bad.append("段落里没有 run")
    pPr = paragraph._p.find(qn("w:pPr"))
    ind = pPr.find(qn("w:ind")) if pPr is not None else None
    indent = spec.first_line_indent
    if getattr(indent, "kind", None) is not None and indent.kind.name == "CHARS":
        chars = ind.get(qn("w:firstLineChars")) if ind is not None else None
        want = str(int(round(indent.chars * 100)))
        if chars != want:
            bad.append(f"首行缩进 firstLineChars={chars!r} ≠ {want!r}")
    return bad


def case_appendix_entries() -> None:
    """S4：`附录A` 这类段落要套 S20（附录序号题名）。"""
    spec = load_preset(PRESET).styles.get("S20")
    assert spec is not None
    source = WORK / "appendix.docx"
    _write(source, [
        ("论测试文档里标题的写法", None),
        ("", None),
        ("附录", None),
        ("附录正文示例。", None),
        ("附录A　调查问卷", None),
        ("问题一：……", None),
    ])
    product = _run(source, "appendix")
    document = Document(str(product))
    target = next((p for p in document.paragraphs
                   if p.text.strip().startswith("附录A")), None)
    check("附录序号题名段落还在", target is not None,
          "产物里找不到『附录A…』段", "S4")
    if target is not None:
        bad = _s20_mismatches(target, spec)
        check("附录A 段落带着 S20 的格式（左对齐/宋体五号/缩进 2 字符）",
              not bad, "；".join(bad), "S4")


def case_list_fields() -> None:
    """清单域插入：有题注 → 有域；没题注 → 不插空域。"""
    code_figure = '图题,1'
    caption_doc = WORK / "caption.docx"
    _write(caption_doc, [
        ("论测试文档里标题的写法", None),
        ("", None),
        ("图 1 测试图题", "Caption"),
        ("正文一段。", None),
    ])
    product = _run(caption_doc, "caption", enable={"插图清单": True})
    codes = _instr_codes(product)
    check("有图题时插入插图清单域（\\t 种）",
          any(code_figure in code for code in codes),
          f"块指令={codes[:3]!r}", "清单域·sections._LIST_FIELD_CODE")

    plain_doc = WORK / "plain.docx"
    _write(plain_doc, [
        ("论测试文档里标题的写法", None),
        ("", None),
        ("正文一段，没有任何题注。", None),
    ])
    product = _run(plain_doc, "plain", enable={"插图清单": True})
    codes = _instr_codes(product)
    check("没有图题时不插空插图清单域",
          not any(code_figure in code for code in codes),
          f"块指令={codes[:3]!r}", "清单域·空清单不插")


# ---------------------------------------------------------------- 2026-09-20 新增
# 逻辑表总则一/三/四：直接改预设开关 → 跑管线 → 查成品。


def _no_structure(config) -> None:
    """结构要求全关——夹具只测自己那一段，不走\"缺件补占位语\"。"""
    config.structure_required = {k: False for k in config.structure_required}


def _first_with(path, needle):
    for paragraph in Document(str(path)).paragraphs:
        if needle in paragraph.text:
            return paragraph
    return None


def _effective_break(paragraph):
    """段落上没写显式值时，段前分页跟随它挂的样式。"""
    direct = paragraph.paragraph_format.page_break_before
    if direct is not None:
        return direct
    return paragraph.style.paragraph_format.page_break_before


def case_label_page_break() -> None:
    """总则四：标签前\"另起一页\"覆盖样式；没写跟随样式；补件插出来的也生效。"""
    specs = load_preset(PRESET).styles
    s07 = specs["S07"].page_break_before     # 摘要标签：预设里不带段前分页
    s02 = specs["S02"].page_break_before     # 参考文献标签：预设里带段前分页
    check("前置：预设里 S07 段前分页=False、S02=True（不然这组没得比）",
          s07 is False and s02 is True, f"S07={s07!r}、S02={s02!r}", "总则四")

    source = WORK / "labels.docx"
    document = Document()
    document.add_paragraph("论测试文档里标题的写法")
    document.add_paragraph("摘要")
    document.add_paragraph("这是摘要正文，供测试使用。")
    document.add_paragraph("一、测试标题")
    document.add_paragraph("参考文献")
    document.add_paragraph("张三. 论文[J]. 2020.")
    document.save(str(source))

    def overrides(config):
        _no_structure(config)
        config.label_page_break = {"摘要": True, "参考文献": False}

    product = _run(source, "labels-override", tweak=overrides)
    label = _first_with(product, "摘要")
    check("摘要标签：覆盖 True（样式是 False）→ 段前分页 True",
          label is not None and label.paragraph_format.page_break_before is True,
          f"取到={None if label is None else label.paragraph_format.page_break_before!r}",
          "总则四")
    label = _first_with(product, "参考文献")
    check("参考文献标签：覆盖 False（样式是 True）→ 段前分页 False",
          label is not None and label.paragraph_format.page_break_before is False,
          f"取到={None if label is None else label.paragraph_format.page_break_before!r}",
          "总则四")

    product = _run(source, "labels-follow", tweak=_no_structure)
    label = _first_with(product, "摘要")
    check("摘要标签：没写覆盖 → 跟随样式（False）",
          label is not None and _effective_break(label) is False,
          f"有效值={None if label is None else _effective_break(label)!r}", "总则四")
    label = _first_with(product, "参考文献")
    check("参考文献标签：没写覆盖 → 段落上无显式值、样式 True 照旧",
          label is not None
          and label.paragraph_format.page_break_before is None
          and _effective_break(label) is True,
          f"显式={None if label is None else label.paragraph_format.page_break_before!r}"
          f" 有效={None if label is None else _effective_break(label)!r}", "总则四")

    source2 = WORK / "labels-generated.docx"
    document = Document()
    document.add_paragraph("论测试文档里标题的写法")
    document.add_paragraph("一、测试标题")
    document.add_paragraph("正文内容，供测试使用。")
    document.save(str(source2))

    def generate(config):
        _no_structure(config)
        config.structure_required["摘要"] = True
        config.label_page_break = {"摘要": True}

    label = _first_with(_run(source2, "labels-generated", tweak=generate), "摘要")
    check("缺件补出来的摘要标签：覆盖 True 也生效（生成路径）",
          label is not None and label.paragraph_format.page_break_before is True,
          f"取到={'无' if label is None else label.paragraph_format.page_break_before!r}",
          "总则四")


def _body_seq(path) -> list[tuple[str, str, int]]:
    """成品 body 顶层顺序：`[(标签, 文字, 手打分页符数)]`（p/tbl，忽略节定义）。"""
    document = Document(str(path))
    seq: list[tuple[str, str, int]] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag not in ("p", "tbl"):
            continue
        text = "".join(t.text or "" for t in child.iter(qn("w:t")))
        breaks = sum(1 for br in child.iter(qn("w:br"))
                     if br.get(qn("w:type")) == "page")
        seq.append((tag, text, breaks))
    return seq


def case_cleanup() -> None:
    """总则三：清手打分页符/空行——封面除外；表格之间与文末表格之后留空段。"""
    source = WORK / "cleanup.docx"
    document = Document()
    cover = document.add_paragraph("封面标题")
    cover.add_run().add_break(WD_BREAK.PAGE)     # 封面里的手打分页符：不该被清
    document.add_paragraph("")                   # 封面里的空行：不该被清
    document.add_paragraph("正文段落甲。")
    document.add_paragraph("　　　")             # 全角空格：不算空行
    only_break = document.add_paragraph("")      # 只含手打分页符的段落
    only_break.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表甲"
    document.add_paragraph("")                   # 两张表格之间：要留
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表乙"
    document.add_paragraph("")                   # 文末表格之后：要留
    document.save(str(source))

    def both(config):
        _no_structure(config)
        config.clear_page_breaks = True
        config.clear_blank_lines = True

    seq = _body_seq(_run(source, "cleanup-both", tweak=both, cover_end=1))
    check("封面里的手打分页符保住、空行也保住",
          seq[0][2] == 1 and seq[1] == ("p", "", 0),
          f"前两段={seq[:2]!r}", "总则三·封面除外")
    check("正文里的手打分页符清掉（全篇只剩封面那一个）",
          sum(item[2] for item in seq) == 1,
          f"每段分页符数={[item[2] for item in seq]!r}", "总则三")
    texts = [item[1] for item in seq if item[0] == "p"]
    check("双勾选后正文形态：甲→全角空格、只含分页符的段一并清掉",
          texts == ["封面标题", "", "正文段落甲。", "　　　", "", ""],
          f"段落={texts!r}", "总则三")
    tables = [i for i, item in enumerate(seq) if item[0] == "tbl"]
    check("两张表格之间留一个空段",
          tables and seq[tables[0] + 1] == ("p", "", 0),
          f"表间={seq[tables[0] + 1] if tables else '无表'!r}", "总则三·防并表")
    check("文末表格之后留一个空段",
          seq[-1] == ("p", "", 0), f"末尾={seq[-1]!r}", "总则三·正文收尾")

    def blank_only(config):
        _no_structure(config)
        config.clear_blank_lines = True

    seq = _body_seq(_run(source, "cleanup-blank", tweak=blank_only, cover_end=1))
    check("只勾“清除空行”：只含分页符的段落不算空、原样保留",
          ("p", "", 1) in seq and sum(item[2] for item in seq) == 2,
          f"段序列={seq!r}", "总则三")
    texts = [item[1] for item in seq if item[0] == "p"]
    check("只勾“清除空行”：正文空段照清、封面空行照留",
          texts[:4] == ["封面标题", "", "正文段落甲。", "　　　"],
          f"段落={texts!r}", "总则三")

    seq = _body_seq(_run(source, "cleanup-none", tweak=_no_structure, cover_end=1))
    empty = [item for item in seq if item[0] == "p" and item[1] == ""]
    check("两个开关都不勾：空段与分页符原样保留（4 空段 + 1 分页符段）",
          len(empty) == 5 and sum(item[2] for item in seq) == 2,
          f"空段={len(empty)} 分页符段={sum(1 for i in seq if i[2])}", "总则三·默认关")


def _run_of(path, para_text, run_text):
    for paragraph in Document(str(path)).paragraphs:
        if para_text in paragraph.text:
            for run in paragraph.runs:
                if run_text in run.text:
                    return run
    return None


def _flag(run, tag) -> bool | None:
    """run 上的开关元素：None=没写；True/False=写了（val=0/false 算 False）。"""
    if run is None:
        return None
    rPr = run._element.find(qn("w:rPr"))
    element = rPr.find(qn(tag)) if rPr is not None else None
    if element is None:
        return None
    return element.get(qn("w:val")) not in ("0", "false")


def _color(run) -> str | None:
    if run is None:
        return None
    rPr = run._element.find(qn("w:rPr"))
    element = rPr.find(qn("w:color")) if rPr is not None else None
    return None if element is None else (element.get(qn("w:val")) or "").upper()


def case_keep_run_formatting() -> None:
    """总则一：清直接格式时保留哪些字符格式（2026-09-20 增\"颜色\"）。"""
    source = WORK / "keep.docx"
    document = Document()
    document.add_paragraph("论测试文档里标题的写法")
    heading = document.add_paragraph()
    run = heading.add_run("一、测试标题")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)
    body = document.add_paragraph()
    run = body.add_run("加粗内容"); run.font.bold = True
    run = body.add_run("彩色内容"); run.font.color.rgb = RGBColor(0xFF, 0, 0)
    run = body.add_run("斜体内容"); run.font.italic = True
    run = body.add_run("下划线内容"); run.font.underline = True
    run = body.add_run("双料内容"); run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0xFF)
    document.add_paragraph("参考文献")
    entry = document.add_paragraph()
    run = entry.add_run("张三. 论文[J]. 2020. ")
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xAA, 0, 0)
    run = entry.add_run("https://example.com")
    run.font.color.rgb = RGBColor(0, 0, 0xAA)
    run.font.underline = True
    document.save(str(source))

    product = _run(source, "keep-default", tweak=_no_structure)
    bold = _run_of(product, "加粗内容", "加粗内容")
    italic = _run_of(product, "斜体内容", "斜体内容")
    under = _run_of(product, "下划线内容", "下划线内容")
    color = _run_of(product, "彩色内容", "彩色内容")
    both = _run_of(product, "双料内容", "双料内容")
    check("默认保留：加粗/斜体/下划线三样都留着",
          _flag(bold, "w:b") is True and _flag(italic, "w:i") is True
          and _flag(under, "w:u") is True,
          f"b={_flag(bold, 'w:b')!r} i={_flag(italic, 'w:i')!r} u={_flag(under, 'w:u')!r}",
          "总则一")
    check("默认保留：颜色清掉（双料内容只剩加粗）",
          _color(color) is None and _flag(both, "w:b") is True
          and _color(both) is None,
          f"彩色={_color(color)!r} 双料color={_color(both)!r}", "总则一")
    head_run = _run_of(product, "测试标题", "测试标题")
    check("一至四级标题连加粗/颜色都不留",
          _flag(head_run, "w:b") is not True and _color(head_run) is None,
          f"b={_flag(head_run, 'w:b')!r} color={_color(head_run)!r}", "总则一·标题例外")

    def with_color(config):
        _no_structure(config)
        config.keep_run_formatting = frozenset({"b", "i", "u", "vertAlign", "color"})

    product = _run(source, "keep-color", tweak=with_color)
    color = _run_of(product, "彩色内容", "彩色内容")
    both = _run_of(product, "双料内容", "双料内容")
    head_run = _run_of(product, "测试标题", "测试标题")
    check("勾上颜色：run 上的颜色留着（FF0000 / 0000FF）",
          _color(color) == "FF0000" and _color(both) == "0000FF",
          f"彩色={_color(color)!r} 双料={_color(both)!r}", "总则一·颜色")
    check("勾上颜色也不动标题：标题仍无颜色、无加粗",
          _flag(head_run, "w:b") is not True and _color(head_run) is None,
          f"b={_flag(head_run, 'w:b')!r} color={_color(head_run)!r}", "总则一·标题例外")

    def clear_all(config):
        _no_structure(config)
        config.keep_run_formatting = frozenset()

    product = _run(source, "keep-none", tweak=clear_all)
    bold = _run_of(product, "加粗内容", "加粗内容")
    under = _run_of(product, "下划线内容", "下划线内容")
    check("“全清”：加粗/下划线一个不留",
          _flag(bold, "w:b") is not True and _flag(under, "w:u") is not True,
          f"b={_flag(bold, 'w:b')!r} u={_flag(under, 'w:u')!r}", "总则一")
    url = _run_of(product, "张三", "example.com")
    plain = _run_of(product, "张三", "张三")
    check("参考文献里的链接文字：颜色与划线始终保留",
          _color(url) == "0000AA" and _flag(url, "w:u") is True,
          f"color={_color(url)!r} u={_flag(url, 'w:u')!r}", "总则一·链接")
    check("同一条目里的普通文字对比组：照常清",
          _color(plain) is None and _flag(plain, "w:b") is not True,
          f"color={_color(plain)!r} b={_flag(plain, 'w:b')!r}", "总则一·链接")


def main() -> int:
    started = time.perf_counter()
    shutil.rmtree(WORK, ignore_errors=True)
    print("== S4 附录序号题名 ==")
    case_appendix_entries()
    print("== 清单域插入计划 ==")
    case_list_fields()
    print("== 标签前另起一页（总则四）==")
    case_label_page_break()
    print("== 清理：手打分页符/空行（总则三）==")
    case_cleanup()
    print("== 保留的字符格式（总则一）==")
    case_keep_run_formatting()
    print()
    print(f"失败 {len(fails)} 条；问题条数：{len(fails)}"
          f"（耗时 {time.perf_counter() - started:.1f}s）")
    REPORT.write_text("\n".join(fails), encoding="utf-8")
    shutil.rmtree(WORK, ignore_errors=True)
    return 1 if fails else 0


if __name__ == "__main__":
    raise SystemExit(main())
