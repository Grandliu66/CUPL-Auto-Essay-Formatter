# -*- coding: utf-8 -*-
r"""Word 实测层：预设「格式」大表格里的每个角色、每个属性，**在 Word 里实际生效的值**是否等于预设。

为什么要有这一层：前面各层看的是 docx 的 XML（样式表里写了什么），而 Word 渲染时还要把
**继承链、行数/自动写法、编号、"自定义样式"标记**这些都算进去——XML 里写对了、Word 里不生效的
情形，只有让真的 Word 打开、读出"实际值"才看得见。2026-10-11 用户反馈的三个问题都属于这一类：
  * 二级标题段前是源文档的"0.5 行"、不是预设的 18 磅（`w:beforeLines` 盖过 `w:before`）；
  * 目录条目是四号 + 首行缩进（新建的 `toc 1/2/3` 带"自定义样式"标记，Word 不认）；
  * （同轮顺带查出）源文档的标题样式挂着多级编号 ⇒ 成品里"第1章 一、引言"双重编号；
    段落上写了"取消编号"（`numId=0`）⇒ 样式里的缩进被 Word 当成 0。

做法：用预设对一批夹具跑完整管线，成品交给本机 Word（COM）打开、刷新域，再逐段读出实际的
字体/字号/加粗/对齐/缩进/段前段后/行距/段前分页，跟预设样式表逐项比。夹具除 reference/ 里的
真实稿子，还有一份**现造的"敌对样式"稿**（段前后按行数/自动、缩进按字符、标题样式挂多级
编号、主题字体、文档网格——全是 Word 里常见的源文档写法）。

⚠️ 需要本机装了 Microsoft Word 和 pywin32；缺任何一样就**跳过并标黄**（不算红，也不算绿）。
默认只跑两套预设 × 5 份夹具（约 2 分钟）；`--full` 跑全部四套预设 × 全部夹具。

用法：
    .venv\Scripts\python.exe "Adversarial Automatic Test\aat_word_effective.py"
    .venv\Scripts\python.exe "Adversarial Automatic Test\aat_word_effective.py" --full
"""
from __future__ import annotations

import collections
import contextlib
import io
import os
import pathlib
import re
import shutil
import sys
import time

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent
for entry in (str(ROOT), str(HERE)):
    if entry not in sys.path:
        sys.path.insert(0, entry)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

WORK = HERE / "_out" / "word_effective"
REPORT = HERE / "_out" / "aat-word-effective.txt"

DEFAULT_PRESETS = ["中国政法大学-本科毕业论文", "中国政法大学-学术十星"]
ALL_PRESETS = ["中国政法大学-本科毕业论文", "中国政法大学-学年论文",
               "中国政法大学-学术十星", "中国政法大学-课程考核"]
DEFAULT_FIXTURES = ["_敌对样式", "空白模板.docx", "民总.docx", "TitleIdentify.docx", "test2.docx"]
ALL_FIXTURES = ["_敌对样式", "空白模板.docx", "民总.docx", "PaperTest.docx", "一般情景测试.docx",
                "图表压力测试.docx", "图表脚注测试.docx", "平时作业.docx", "test2.docx",
                "TitleIdentify.docx", "压力！.docx", "阿拉伯数字.docx"]
PER_ROLE_LIMIT = 8        # 每份夹具里每个角色最多查几段


def _import_word_stack():
    """返回 (win32com.client, 原因)；没有就 (None, 原因)。"""
    try:
        import win32com.client                                   # noqa: F401
        return sys.modules["win32com.client"], ""
    except ImportError:
        pass
    # 项目 .venv / 打包用 venv 里装着 pywin32，但用别的解释器跑时不在 sys.path 上
    for venv in (ROOT / ".venv", ROOT / "packaging" / ".venv"):
        site = venv / "Lib" / "site-packages"
        if not site.is_dir():
            continue
        for sub in ("", "win32", "win32/lib", "pythonwin", "pywin32_system32"):
            path = str((site / sub).resolve())
            if path not in sys.path:
                sys.path.append(path)
            with contextlib.suppress(Exception):
                os.add_dll_directory(path)
        try:
            import win32com.client                               # noqa: F401
            return sys.modules["win32com.client"], ""
        except ImportError:
            continue
    return None, "没有 pywin32"


# ---------------------------------------------------------------- 敌对样式夹具

def build_hostile(path: pathlib.Path) -> None:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn

    doc = Document()
    w = nsdecls("w")

    def set_style(style_id, ppr_frags, rpr_frags=None):
        style = [s for s in doc.styles if s.style_id == style_id][0]
        pPr = style.element.get_or_add_pPr()
        for child in list(pPr):
            pPr.remove(child)
        for frag in ppr_frags:
            pPr.append(parse_xml(frag))
        if rpr_frags is not None:
            rPr = style.element.get_or_add_rPr()
            for child in list(rPr):
                rPr.remove(child)
            for frag in rpr_frags:
                rPr.append(parse_xml(frag))

    numbering = doc.part.numbering_part.element
    abstract = parse_xml(
        f'<w:abstractNum {w} w:abstractNumId="90"><w:multiLevelType w:val="multilevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading1"/>'
        '<w:lvlText w:val="第%1章"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="432" w:hanging="432"/></w:pPr></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading2"/>'
        '<w:lvlText w:val="%1.%2"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="576" w:hanging="576"/></w:pPr></w:lvl>'
        '</w:abstractNum>')
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(parse_xml(f'<w:num {w} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))

    set_style("Normal",
              [f'<w:snapToGrid {w} w:val="1"/>',
               f'<w:spacing {w} w:before="100" w:beforeAutospacing="1" w:after="100" '
               'w:afterAutospacing="1" w:line="360" w:lineRule="auto"/>',
               f'<w:ind {w} w:firstLineChars="300" w:firstLine="630"/>', f'<w:jc {w} w:val="both"/>'],
              [f'<w:rFonts {w} w:asciiTheme="minorHAnsi" w:eastAsiaTheme="minorEastAsia" w:hAnsiTheme="minorHAnsi"/>',
               f'<w:sz {w} w:val="21"/>'])
    set_style("Heading1",
              [f'<w:keepNext {w}/>', f'<w:numPr {w}><w:numId w:val="90"/></w:numPr>',
               f'<w:spacing {w} w:beforeLines="100" w:afterLines="100" w:line="360" w:lineRule="auto"/>',
               f'<w:ind {w} w:leftChars="100" w:hangingChars="200"/>', f'<w:jc {w} w:val="center"/>',
               f'<w:outlineLvl {w} w:val="0"/>'],
              [f'<w:rFonts {w} w:asciiTheme="majorHAnsi" w:eastAsiaTheme="majorEastAsia"/>', f'<w:b {w}/>',
               f'<w:color {w} w:val="2F5496" w:themeColor="accent1" w:themeShade="BF"/>', f'<w:sz {w} w:val="44"/>'])
    set_style("Heading2",
              [f'<w:keepNext {w}/>', f'<w:numPr {w}><w:ilvl w:val="1"/><w:numId w:val="90"/></w:numPr>',
               f'<w:spacing {w} w:beforeLines="50" w:afterAutospacing="1"/>',
               f'<w:ind {w} w:left="567" w:hanging="567"/>', f'<w:outlineLvl {w} w:val="1"/>'],
              [f'<w:b {w}/>', f'<w:sz {w} w:val="32"/>'])
    set_style("Heading3",
              [f'<w:contextualSpacing {w}/>', f'<w:spacing {w} w:before="240" w:after="60"/>',
               f'<w:ind {w} w:left="720" w:hanging="360"/>', f'<w:outlineLvl {w} w:val="2"/>'],
              [f'<w:i {w}/>', f'<w:sz {w} w:val="26"/>'])
    set_style("Heading4",
              [f'<w:spacing {w} w:beforeLines="30" w:afterLines="30"/>', f'<w:ind {w} w:leftChars="200"/>',
               f'<w:outlineLvl {w} w:val="3"/>'],
              [f'<w:sz {w} w:val="24"/>'])

    for text, style in [
            ("论文题目", None), ("摘要", None), ("摘要内容。", None), ("关键词：甲；乙", None),
            ("一、引言", "Heading 1"), ("正文第一段，用来测试正文格式。", None),
            ("（一）第一节", "Heading 2"), ("节下正文。", None),
            ("1. 三级标题", "Heading 3"), ("三级下正文。", None),
            ("（1）四级标题", "Heading 4"), ("四级下正文。", None),
            ("二、第二章", "Heading 1"), ("（一）又一节", "Heading 2"), ("正文。", None),
            ("参考文献", None), ("[1] 张三. 民法总论[M]. 北京: 法律出版社, 2020.", None)]:
        doc.add_paragraph(text, style=style) if style else doc.add_paragraph(text)
    # 一段"取消编号"（numId=0）的标题：Word 里点掉编号后留下的就是这个
    doc.add_paragraph("2. 取消了编号的三级标题", style="Heading 3")
    doc.paragraphs[-1]._p.get_or_add_pPr().append(
        parse_xml(f'<w:numPr {w}><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr>'))
    doc.sections[0]._sectPr.append(parse_xml(f'<w:docGrid {w} w:type="lines" w:linePitch="312"/>'))
    doc.save(str(path))


# ---------------------------------------------------------------- 比对

def _check_paragraph(spec, fmt, page_break_expect=None):
    """段落级属性；返回 [(属性, 期望, 实际)]。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from essay_tool.config.units import IndentKind

    align = {WD_ALIGN_PARAGRAPH.LEFT: 0, WD_ALIGN_PARAGRAPH.CENTER: 1,
             WD_ALIGN_PARAGRAPH.RIGHT: 2, WD_ALIGN_PARAGRAPH.JUSTIFY: 3}

    def close(a, b, tol=0.06):
        return abs(a - b) <= tol

    bad = []
    if fmt.Alignment != align[spec.alignment]:
        bad.append(("对齐", align[spec.alignment], fmt.Alignment))
    for which, ind, pt_attr, chars_attr in (
            ("首行缩进", spec.first_line_indent, "FirstLineIndent", "CharacterUnitFirstLineIndent"),
            ("左缩进", spec.left_indent, "LeftIndent", "CharacterUnitLeftIndent")):
        pts, chars = getattr(fmt, pt_attr), getattr(fmt, chars_attr)
        if ind.kind == IndentKind.NONE:
            if abs(pts) > 0.06 or abs(chars) > 0.01:
                bad.append((which, "0", f"{pts}pt/{chars}字符"))
        elif ind.kind == IndentKind.CHARS:
            if not close(chars, ind.chars, 0.05):
                bad.append((which, f"{ind.chars}字符", f"{pts}pt/{chars}字符"))
        elif not close(pts, ind.length.pt, 0.3):
            bad.append((which, f"{ind.length.pt:.1f}pt", f"{pts}pt/{chars}字符"))
    for which, value, attr, lines, auto in (
            ("段前", spec.space_before, "SpaceBefore", "LineUnitBefore", "SpaceBeforeAuto"),
            ("段后", spec.space_after, "SpaceAfter", "LineUnitAfter", "SpaceAfterAuto")):
        if value is None:
            continue                                   # 预设留空＝继承，不查
        if not close(getattr(fmt, attr), value.pt) or getattr(fmt, lines) != 0 or getattr(fmt, auto):
            bad.append((which, f"{value.pt}pt",
                        f"{getattr(fmt, attr)}pt/{getattr(fmt, lines)}行/自动={bool(getattr(fmt, auto))}"))
    rule = {WD_LINE_SPACING.SINGLE: (0, None), WD_LINE_SPACING.ONE_POINT_FIVE: (1, None),
            WD_LINE_SPACING.DOUBLE: (2, None), WD_LINE_SPACING.AT_LEAST: (3, spec.line_value),
            WD_LINE_SPACING.EXACTLY: (4, spec.line_value), WD_LINE_SPACING.MULTIPLE: (5, None)}.get(spec.line_rule)
    if rule and rule[0] == 5:
        # Word 把"多倍 1/1.5/2"报成单倍/1.5 倍/双倍，是同一回事
        rule = {1.0: (0, None), 1.5: (1, None), 2.0: (2, None)}.get(float(spec.line_value), rule)
    if rule:
        if fmt.LineSpacingRule != rule[0]:
            bad.append(("行距类型", rule[0], fmt.LineSpacingRule))
        elif rule[0] in (3, 4) and not close(fmt.LineSpacing, rule[1].pt):
            bad.append(("行距值", f"{rule[1].pt}pt", fmt.LineSpacing))
        elif rule[0] == 5 and not close(fmt.LineSpacing, 12 * float(spec.line_value), 0.2):
            bad.append(("行距倍数", spec.line_value, fmt.LineSpacing / 12))
    expect_break = spec.page_break_before if page_break_expect is None else page_break_expect
    if bool(fmt.PageBreakBefore) != bool(expect_break):
        bad.append(("段前分页", expect_break, bool(fmt.PageBreakBefore)))
    return bad


def _check_font(spec, char_range, char):
    bad = []
    font = char_range.Font
    if abs(font.Size - spec.size.pt) > 0.05:
        bad.append(("字号", spec.size.pt, font.Size))
    if bool(font.Bold) != bool(spec.bold):
        bad.append(("加粗", spec.bold, bool(font.Bold)))
    if re.match("[一-鿿　-〿＀-￯]", char):
        if font.NameFarEast != spec.font_east_asian:
            bad.append(("中文字体", spec.font_east_asian, font.NameFarEast))
    elif char.isascii() and char.isalpha():
        if font.NameAscii != spec.font_ascii:
            bad.append(("西文字体", spec.font_ascii, font.NameAscii))
    return bad


def _roles(entry, first_appendix_content):
    """识别层的身份 → 它该套预设的哪个角色：(段落样式,) 或同段关键词的 (段落样式, 标签样式, 内容样式)。"""
    from essay_tool.engine import apply_structure as A

    part, slot, meta = entry.part, entry.meta.get("slot"), entry.meta
    levels = {"一级标题": "S02", "二级标题": "S03", "三级标题": "S04", "四级标题": "S05"}
    fixed = {"题目": "S06", "副标题": "S29", "英文题目": "S11", "英文副标题": "S32"}
    labels = A._LABEL_CONTENT_STYLE_IDS
    if part in fixed:
        return (fixed[part],)
    if part in levels:
        return (levels[part],)
    if part in ("关键词", "英文关键词") and meta.get("inline"):
        return (labels[part][0], labels[part][0], labels[part][1])
    if part in labels and slot == "label":
        return (labels[part][0],)
    if part in labels and slot == "content" and labels[part][1]:
        return (labels[part][1],)
    if part in A._ANCHOR_STYLE_IDS and slot == "label":
        return (A._ANCHOR_STYLE_IDS[part],)
    if part == "附录" and slot == "content":
        return ("S20",) if first_appendix_content else ("S01",)
    if part in A._ANCHOR_CONTENT_STYLE_IDS and slot == "content":
        return (A._ANCHOR_CONTENT_STYLE_IDS[part],)
    if not part:
        return ("S01",)
    return None


def audit_one(word, preset, preset_name, source, tag, problems):
    """跑一份夹具，返回检查了多少段。"""
    from docx import Document
    from essay_tool.engine import pipeline
    from essay_tool.engine.doc_index import paragraph_ordinals
    from essay_tool.engine.recognize import recognize

    override = {key: True for key in preset.structure_required}      # 所有部件都开，各角色都出现
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            out = pipeline.run(source, None, WORK / f"{preset_name}-{tag}", preset_config=preset,
                               structure_override=override)
    except Exception as exc:                                           # noqa: BLE001
        problems.append(f"{preset_name} · {tag}：管线失败 {exc!r}"[:160])
        return 0

    document = Document(str(out))
    index = recognize(document)
    ordinals = paragraph_ordinals(index)
    expected = []
    appendix_seen = False
    for i, paragraph in enumerate(document.paragraphs):
        entry = ordinals.get(i)
        text = paragraph.text.strip()
        if entry is None or not text:
            continue
        caption = entry.meta.get("caption")
        if caption:
            if caption.get("role") == "main" and caption.get("kind") in ("图", "表"):
                expected.append((text, ("S24" if caption["kind"] == "图" else "S25",), entry))
            continue
        if paragraph.style is not None and not entry.part and paragraph.style.style_id != "Normal":
            continue
        if (not entry.part or (entry.part == "附录" and entry.meta.get("slot") == "content")) \
                and re.match(r"^附\s*录\s*[A-Za-z]", text):
            expected.append((text, ("S20",), entry))                   # 附录序号题名：引擎另按 S20 套
            if entry.part == "附录":
                appendix_seen = True                                   # 它就是"附录第一段"，之后的是 S01
            continue
        first = False
        if entry.part == "附录" and entry.meta.get("slot") == "content" and not appendix_seen:
            first, appendix_seen = True, True
        roles = _roles(entry, first)
        if roles is None:
            continue
        if not entry.part and paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
            continue                                                   # 正文里原有的编号列表，不归预设管
        expected.append((text, roles, entry))

    try:
        wd = word.Documents.Open(os.path.normpath(str(out)), ReadOnly=False)
    except Exception as exc:                                           # noqa: BLE001
        print(f"   · {preset_name} · {tag}：Word 打不开成品，跳过（{str(exc)[-30:]}）")
        return 0
    checked = 0
    try:
        wd.Fields.Update()
        wparas = [(p.Range.Text.strip().replace("\x07", ""), p) for p in wd.Paragraphs]
        pointer = 0
        counts = collections.Counter()
        skip_styles = ("toc", "caption", "题注", "图题", "表题", "table of")

        def report(sid, bad, text):
            for prop, want, got in bad:
                problems.append(f"{preset_name} · {tag} · {sid} {preset.styles[sid].name} · {prop}："
                                f"期望 {want!r} 实际 {got!r}（「{text[:14]}」）")

        for text, roles, entry in expected:
            j = pointer
            while j < len(wparas) and (wparas[j][0] != text or (
                    roles[0] not in ("S24", "S25")
                    and wparas[j][1].Style.NameLocal.lower().startswith(skip_styles))):
                j += 1
            if j >= len(wparas):
                continue
            pointer = j + 1
            para = wparas[j][1]
            if para.Range.Information(12):                             # 在表格里，另查
                continue
            sid = roles[0]
            if counts[sid] >= PER_ROLE_LIMIT:
                continue
            counts[sid] += 1
            checked += 1
            spec = preset.styles[sid]
            override_break = None
            if entry.part in preset.label_page_break and entry.meta.get("slot") == "label":
                override_break = preset.label_page_break[entry.part]
            bad = _check_paragraph(spec, para.Format, override_break)
            rng = para.Range
            last = max(rng.Characters.Count - 1, 1)
            if sid != "S21" and rng.ListFormat.ListType != 0:
                bad.append(("意外的自动编号", "无", rng.ListFormat.ListString))
            if len(roles) == 3:                                        # 同段关键词：标签一截、关键词一截
                bad += _check_font(preset.styles[roles[1]], rng.Characters(1), text[0])
                bad += [("[关键词部分]" + a, b, c) for a, b, c in
                        _check_font(preset.styles[roles[2]], rng.Characters(last), text[-1])]
            else:
                bad += _check_font(spec, rng.Characters(1), text[0])
                if len(text) > 2:
                    bad += [("[末字]" + a, b, c) for a, b, c in _check_font(spec, rng.Characters(last), text[-1])]
            report(sid, bad, text)

        # 目录条目 → S15~S17（Word 刷新目录域后读）
        for text, para in wparas:
            m = re.match(r"toc (\d)$", para.Style.NameLocal.lower())
            if m and text:
                sid = {"1": "S15", "2": "S16", "3": "S17"}.get(m.group(1))
                if sid and counts[sid] < 4:
                    counts[sid] += 1
                    checked += 1
                    report(sid, _check_paragraph(preset.styles[sid], para.Format)
                           + _check_font(preset.styles[sid], para.Range.Characters(1), text[0]), text)
        # 表格文字 → S26
        for ti in range(1, min(wd.Tables.Count, 3) + 1):
            seen = 0
            for cell in wd.Tables(ti).Range.Cells:
                text = cell.Range.Text.strip().replace("\x07", "").replace("\r", "")
                if not text or seen >= 3:
                    continue
                seen += 1
                cp = cell.Range.Paragraphs(1)
                report("S26", _check_paragraph(preset.styles["S26"], cp.Format)
                       + _check_font(preset.styles["S26"], cp.Range.Characters(1), text[0]), text)
            checked += seen
        # 脚注文本 → S23
        if wd.Footnotes.Count:
            fp = wd.Footnotes(1).Range.Paragraphs(1)
            report("S23", _check_paragraph(preset.styles["S23"], fp.Format), fp.Range.Text.strip())
            checked += 1
    finally:
        wd.Close(False)
    return checked


def main() -> int:
    global word
    started = time.perf_counter()
    full = "--full" in sys.argv
    win32, reason = _import_word_stack()
    if win32 is None:
        print(f"⚠️ 跳过：{reason}，Word 实测层没有跑。")
        return 0
    try:
        word = win32.DispatchEx("Word.Application")
    except Exception as exc:                                           # noqa: BLE001
        print(f"⚠️ 跳过：本机启动不了 Word（{str(exc)[:60]}），Word 实测层没有跑。")
        return 0
    word.Visible = False
    word.DisplayAlerts = 0

    from essay_tool.config.preset_io import load_preset

    shutil.rmtree(WORK, ignore_errors=True)
    WORK.mkdir(parents=True, exist_ok=True)
    problems: list[str] = []
    total = 0
    try:
        sources = []
        for name in (ALL_FIXTURES if full else DEFAULT_FIXTURES):
            if name == "_敌对样式":
                path = WORK / "_敌对样式.docx"
                build_hostile(path)
            else:
                path = ROOT / "reference" / name
                if not path.exists():
                    print(f"   · 夹具缺失，跳过：{name}（reference/ 里没有）")
                    continue
            sources.append((path, path.stem))
        for preset_name in (ALL_PRESETS if full else DEFAULT_PRESETS):
            preset = load_preset(ROOT / "presets" / f"{preset_name}.json")
            for path, tag in sources:
                checked = audit_one(word, preset, preset_name, path, tag, problems)
                total += checked
                print(f"   {preset_name} · {tag}：查了 {checked} 处", flush=True)
    finally:
        word.Quit()

    grouped = collections.Counter(re.sub(r"（「.*?」）$", "", p) for p in problems)
    print()
    for line, n in sorted(grouped.items()):
        print(f"   ✘ {line}" + (f"  ×{n}" if n > 1 else ""))
    print(f"\n共查 {total} 处；问题条数：{len(problems)}（耗时 {time.perf_counter() - started:.1f}s）")
    REPORT.write_text("\n".join(problems), encoding="utf-8")
    return 1 if problems else 0


word = None

if __name__ == "__main__":
    raise SystemExit(main())
